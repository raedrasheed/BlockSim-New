#!/usr/bin/env python3
"""Steps C-I: renumber Ch3 headings, insert new sections/table (RED), append
red paragraphs, add LoT entry, mark fields dirty. Operates on the A/B output."""
import copy
import docx
from docx.shared import RGBColor, Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

FILE = "Raed-Rasheed-draft-37-Related-Work-Added.docx"
RED = RGBColor(0xFF, 0x00, 0x00)
d = docx.Document(FILE)
P = d.paragraphs

def redden(r): r.font.color.rgb = RED

# ---- capture anchors by index (unchanged since A/B added only table rows) ----
A_ch3   = P[340]   # old '3.7. Comparative Analysis...' (new 3.7 + table go before it)
A_57    = P[586]   # '5.7. PoCol Genesis Round'
A_512   = P[737]   # '5.12. Summary'
A_45    = P[447]   # '4.5. Goals and Needs for Design'
A_65    = P[773]   # '6.5. Logged Outputs...'
A_75    = P[859]   # '7.5. Sensitivity Analysis'
A_lot   = P[51]    # LoT entry 'Table 6.1...'
# heading paragraphs to renumber (retain colour, change number text only)
RENUM = [(340,"3.7.","3.8."),(342,"3.7.1.","3.8.1."),(347,"3.7.1.1.","3.8.1.1."),
         (349,"3.7.1.2.","3.8.1.2."),(353,"3.7.2.","3.8.2."),(360,"3.7.3.","3.8.3."),
         (378,"3.8.","3.9.")]
renum_paras = [(P[i],o,n) for i,o,n in RENUM]

# ---- STEP C: renumber existing Ch3 headings (retain original colour) ----
for p,old,new in renum_paras:
    done=False
    for r in p.runs:
        if r.text.startswith(old):
            r.text = new + r.text[len(old):]; done=True; break
        if old in r.text and not done:
            r.text = r.text.replace(old,new,1); done=True; break
    assert done, f"renumber failed for {old} in {p.text[:40]!r}"

# ---- helpers to insert red paragraphs before an anchor ----
def ins(anchor, text, style="Normal", center=False, bold=False):
    p = anchor.insert_paragraph_before(text, style=style)
    if center: p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for r in p.runs:
        redden(r)
        if bold: r.font.bold = True
    return p

H2,H3,H4="Heading 2","Heading 3","Heading 4"

# ============================ STEP D: new Ch3 §3.7 =========================
blocks = [
 ("3.7. Parallel, Collaborative, and Energy-Aware Proof-of-Work Approaches",H2,0),
 ("3.7.1. Parallel and Partitioned Mining",H3,0),
 ("Hazari and Mahmoud proposed Parallel Proof-of-Work as an approach for distributing mining computation among cooperating participants through a coordinating manager [133]. The approach retains the conventional hash-target Proof-of-Work puzzle while allocating mining tasks among multiple miners to improve block-generation performance and reduce duplication of assigned work. It also considers coordination and reward allocation within the parallel mining structure.","Normal",0),
 ("Parallel Proof-of-Work is the closest architectural predecessor to PoCol because both approaches divide mining work among multiple participants rather than allowing every participant to perform an entirely uncoordinated search. However, PoCol defines the distributed work more narrowly and binds it explicitly to the consensus round. In PoCol, all participating miners use one identified immutable candidate-block template, and the permitted nonce domain associated with that template is deterministically partitioned into mutually disjoint ranges. Thus, honest miners do not evaluate identical serialized header–nonce inputs within the same collaborative round.","Normal",0),
 ("PoCol therefore does not claim to be the first approach to parallelize Proof-of-Work or distribute mining computation. Its proposed distinction is the integration of deterministic nonce-range separation, immutable common-template agreement, collaborative reward allocation, and explicit post-range operating policies.","Normal",0),
 ("3.7.2. Collaborative and Team-Based Proof-of-Work",H3,0),
 ("StrongChain introduces weak Proof-of-Work headers that make non-winning computational contributions visible and rewardable [134]. Instead of preventing overlapping work before mining, StrongChain recognizes useful evidence of mining effort after it has been generated. Its main contribution lies in reducing reward variance and making mining participation more transparent.","Normal",0),
 ("Collaborative Proof-of-Work forms dynamic mining groups based on participant capabilities and emphasizes secure contribution verification and fair reward distribution [135]. This is directly relevant to PoCol because both protocols seek to reward miners participating in collaborative block production. Collaborative Proof-of-Work, however, focuses primarily on group formation, contribution verification, and reward fairness, while PoCol focuses on deterministic separation of the nonce ranges searched under a common immutable candidate-block template.","Normal",0),
 ("Proof of Team Sprint is a recent team-based collaborative consensus algorithm intended to reduce the energy inefficiency of conventional Proof-of-Work [136]. It organizes miners into teams that collaborate in solving the cryptographic mining puzzle and evaluates the effects of cooperation on energy use and reward fairness. Proof of Team Sprint is therefore the most recent directly related peer-reviewed collaborative consensus approach considered in this thesis.","Normal",0),
 ("PoCol differs from these approaches by defining collaboration at the level of exact nonce-range assignments associated with a common template. Nevertheless, StrongChain, Collaborative Proof-of-Work, and Proof of Team Sprint provide important precedents for contribution visibility, reward fairness, and organized teamwork. These studies also expose a requirement that PoCol must address more rigorously: a miner should not receive a collaborative reward merely by registering or claiming that its assigned range was completed.","Normal",0),
 ("3.7.3. Energy-Aware Mining Policies",H3,0),
 ("Green-PoW reduces energy consumption by restricting mining participation during selected rounds to a smaller group of eligible miners [15]. Its energy-saving mechanism is based primarily on reducing the number of simultaneously active miners. Proof of Team Sprint follows a team-oriented approach in which collaborative execution is intended to reduce the amount of active computation attributed to individual participants [136].","Normal",0),
 ("PoCol adopts a different energy-management principle. At the beginning of a collaborative round, all eligible miners may participate and receive deterministic, mutually disjoint nonce ranges. Under the proposed energy-saving policy, a miner that exhausts its assigned range without finding a valid candidate header stops active hashing and enters a low-power idle state until the deadline of the current collaborative mining round.","Normal",0),
 ("The energy consumed by miner i under this policy is expressed as:","Normal",0),
 ("Eᵢ = Pᵢ(active)·tᵢ(active) + Pᵢ(idle)·tᵢ(idle) + Eᵢ(coord),","Normal",1),
 ("where Pᵢ(active) is active mining power, Pᵢ(idle) is idle power, and Eᵢ(coord) represents communication, coordination, and state-transition overhead. Accordingly, nonce-range partitioning alone does not guarantee energy reduction. Energy savings arise only when miners spend a non-zero period in a power state that consumes less energy than active hashing.","Normal",0),
 ("PoCol may alternatively operate under a continuous-performance policy. In this mode, miners that complete their assigned ranges continue processing newly assigned non-overlapping work rather than entering the idle state. If aggregate hash rate, hardware efficiency, and experiment duration are equal to those of a continuously operating PoW baseline, this mode is expected to consume approximately the same energy. It may, however, improve accepted-block production or transaction throughput under particular evaluated configurations. Such performance improvement must be established experimentally and must not be treated as an automatic consequence of nonce partitioning.","Normal",0),
 ("Green-PoW and PoCol therefore reduce active mining through different mechanisms. Green-PoW reduces the number of miners allowed to participate, whereas PoCol reduces the active hashing duration of miners after they complete their assigned work. This distinction allows PoCol to preserve wider initial participation while exposing an explicit trade-off between lower energy use and continued mining performance.","Normal",0),
 ("3.7.4. Auditable Mining Contribution",H3,0),
 ("A central challenge in collaborative and partitioned mining is verifying that a miner actually performed its assigned computation. A dishonest miner could claim that it exhausted its range, enter the idle state, and later request a share of the collaborative reward without having evaluated the assigned candidate headers.","Normal",0),
 ("Auditable Proof-of-Work provides a mechanism through which miners can produce probabilistic evidence that specified regions of the nonce domain were searched [137]. Although APoW is primarily directed toward detecting work omission and block-withholding behavior in mining pools, its auditing principle is relevant to PoCol. APoW is available as a 2026 arXiv preprint rather than a peer-reviewed journal or conference publication.","Normal",0),
 ("A strengthened PoCol design could associate each assigned nonce range with cryptographic commitments, sampled shares, or auditable search evidence. Such mechanisms would help distinguish miners that genuinely completed their assignments from miners that falsely claimed completion, thereby reducing free-riding, false idling claims, and reward manipulation.","Normal",0),
 ("3.7.5. Comparison with PoCol and Identified Research Gap",H3,0),
 ("The reviewed approaches address different parts of the problem targeted by PoCol. Parallel Proof-of-Work distributes mining computation and provides the closest architectural precedent for parallel work allocation [133]. StrongChain improves contribution visibility by recording weak Proof-of-Work solutions [134]. Collaborative Proof-of-Work emphasizes dynamic group formation, contribution verification, and reward fairness [135]. Proof of Team Sprint organizes miners into collaborative teams and connects teamwork with energy efficiency and fairness [136]. Green-PoW reduces energy consumption by limiting the number of active miners in selected rounds [15]. APoW provides a possible mechanism for proving that assigned nonce regions were actually searched [137].","Normal",0),
 ("However, these approaches do not jointly formulate all of the following elements within one protocol: agreement on one identified immutable candidate-block template, deterministic assignment of mutually disjoint nonce ranges, collaborative reward distribution, explicit post-range idle and continuous operating policies, and verifiable completion of assigned work.","Normal",0),
 ("PoCol is proposed to address this combined design space. Its contribution is not the independent invention of parallel mining, nonce-range distribution, collaborative rewards, or miner idling. Rather, its proposed contribution lies in integrating these mechanisms within one collaborative Proof-of-Work framework and evaluating the resulting energy–performance trade-off.","Normal",0),
]
for text,style,center in blocks:
    ins(A_ch3, text, style=style, center=bool(center))

# caption (red) before anchor (after §3.7 blocks)
ins(A_ch3, "Table 3.9. Comparison of PoCol with parallel, collaborative, auditable, and energy-aware Proof-of-Work approaches", style="Normal", bold=True)

# ---- build comparison table (RED), insert before anchor ----
HEADER=["Characteristic","Parallel PoW","StrongChain","Collaborative PoW","Proof of Team Sprint","Green-PoW","APoW","PoCol"]
ROWS=[
 ["Conventional hash-target puzzle","Yes","Yes","Yes","Yes","Yes","Yes","Yes"],
 ["Parallel or collaborative mining","Yes","Yes","Yes","Yes","Limited by round","Auditing-oriented","Yes"],
 ["Immutable shared candidate template","Coordinated work","Not central","Not central","Not central","Not central","Not central","Explicit"],
 ["Deterministic nonce-range allocation","Work allocation","No","Not central","Not central","No","Audited regions","Yes"],
 ["Explicit mutually disjoint nonce ranges","Implementation-dependent","No","Not central","Not central","No","No","Yes"],
 ["Contribution visibility or verification","Limited","Weak-header evidence","Explicit","Team-based","Runner-up evidence","Probabilistic audit","Requires extension"],
 ["Energy-saving mechanism","Not primary","Not primary","Efficiency-oriented","Team cooperation","Fewer active miners","Not primary","Reduced active time"],
 ["Post-range idle policy","No","No","No","Not explicit","Non-selected miners stop","No","Explicit"],
 ["Continuous-performance policy","Yes","Yes","Yes","Team-dependent","No","No","Explicit"],
 ["Main strength","Parallel work allocation","Contribution visibility","Fair collaboration","Team cooperation","Direct active-miner reduction","Auditable search evidence","Integrated work partitioning and operating policies"],
 ["Main limitation","Manager dependence","Does not prevent overlap","Pool complexity","Team assumptions","Restricted participation","Not a complete consensus protocol","Contribution proof and complete distributed implementation remain incomplete"],
]
tbl = d.add_table(rows=len(ROWS)+1, cols=8)
# copy borders design from an existing comparison table (d.tables[10])
src_tblPr = d.tables[10]._tbl.tblPr
new_tblPr = tbl._tbl.tblPr
for tag in ('w:tblBorders','w:tblStyle','w:tblLayout'):
    el = src_tblPr.find(qn(tag))
    if el is not None:
        ex = new_tblPr.find(qn(tag))
        if ex is not None: new_tblPr.remove(ex)
        new_tblPr.append(copy.deepcopy(el))
def fill(cell,text,bold=False):
    cell.text=""
    p=cell.paragraphs[0]; r=p.add_run(text)
    r.font.color.rgb=RED; r.font.size=Pt(8)
    if bold: r.font.bold=True
for j,h in enumerate(HEADER): fill(tbl.rows[0].cells[j],h,bold=True)
for i,row in enumerate(ROWS,1):
    for j,v in enumerate(row): fill(tbl.rows[i].cells[j],v)
A_ch3._p.addprevious(tbl._tbl)

# note (red) after table, before anchor
ins(A_ch3, "The comparison does not imply that PoCol is universally superior to the examined protocols. Parallel PoW provides the closest architectural baseline, Green-PoW provides the most relevant energy-policy baseline, Proof of Team Sprint provides the most recent peer-reviewed team-based collaborative baseline, and APoW provides a relevant contribution-auditing direction.", style="Normal")

# ============================ STEP F: Ch5 §5.6.6 =========================
ch5=[
 ("5.6.6. Post-Range Operating Policies",H3,0),
 ("PoCol distinguishes nonce-range allocation from the miner’s operating state after completing the assigned search. The protocol may apply either an energy-saving idle policy or a continuous-performance policy. These policies are alternatives and must be identified explicitly for every experimental or deployment configuration.","Normal",0),
 ("5.6.6.1. Energy-Saving Idle Policy",H4,0),
 ("Under the energy-saving policy, a miner searches only the nonce range assigned to it for the current immutable template. If no valid block is found before the range is exhausted, the miner stops active hashing and enters a low-power idle state until the end of the current collaborative mining round. The miner remains connected and may receive protocol messages, including notification that another miner has produced a valid block, but it performs no further mining hashes during the remaining round time.","Normal",0),
 ("The phrase “end of the current collaborative mining round” refers to the round deadline or target block interval defined by PoCol. It does not refer to a multi-block difficulty-adjustment epoch.","Normal",0),
 ("The energy benefit of this policy depends on the difference between active and idle power and on the duration of the idle interval. The policy may also reduce the active network hash rate as miners complete their ranges; therefore, its effects on block interval, range-exhaustion probability, liveness, and security must be evaluated together with its energy savings.","Normal",0),
 ("5.6.6.2. Continuous Performance Policy",H4,0),
 ("Under the continuous-performance policy, a miner that completes its assigned range does not enter the low-power idle state. Instead, it may continue processing newly assigned, non-overlapping work defined by the protocol. The new assignment must be bound to an explicitly identified template or sub-round so that honest miners do not receive overlapping assignments.","Normal",0),
 ("When PoCol and the PoW baseline use the same aggregate hash rate, hardware efficiency, active power, and experiment duration, continuous operation is not expected to provide an inherent energy reduction. Its purpose is performance-oriented: it may improve accepted-block production or transaction throughput under the evaluated configuration. Such improvement must be reported as an empirical result rather than as a guaranteed theoretical consequence of nonce-range separation.","Normal",0),
]
for text,style,center in ch5: ins(A_57,text,style=style,center=bool(center))

# ============================ STEP G: appends =========================
# 5.11.5
for t in [
 "Parallel Proof-of-Work provides the closest architectural comparison to PoCol because both approaches distribute mining computation among cooperating participants [133]. PoCol differs by binding deterministic, mutually disjoint nonce ranges to one identified immutable candidate-block template and by defining explicit miner behavior after an assigned range is exhausted.",
 "Green-PoW provides the most relevant energy-policy comparison [15]. Green-PoW reduces energy use by reducing the number of miners allowed to participate in selected rounds, whereas PoCol’s idle policy initially permits broad participation and subsequently reduces each miner’s active hashing duration after its assigned range is completed.",
 "Proof of Team Sprint provides the most recent peer-reviewed team-based collaborative comparison [136]. Collaborative Proof-of-Work and StrongChain provide relevant precedents for reward fairness, contribution verification, and contribution visibility [135], [134]. APoW provides a potential direction for verifying whether an assigned nonce region was actually searched [137].",
 "These comparisons do not establish universal superiority. PoCol’s proposed distinction is the integration of immutable template agreement, deterministic non-overlapping nonce allocation, collaborative reward distribution, and selectable idle or continuous operating policies. Its complete distributed implementation and contribution-verification mechanism remain areas requiring further development and evaluation.",
]: ins(A_512,t,style="Normal")

# 4.4.2
ins(A_45,"For the PoCol energy-saving policy, total network energy must include both active and idle periods:",style="Normal")
ins(A_45,"E_PoCol,idle = Σᵢ [ Pᵢ(active)·tᵢ(active) + Pᵢ(idle)·tᵢ(idle) + Eᵢ(coord) ].",style="Normal",center=True)
ins(A_45,"For the continuous-performance policy, if all miners remain active for the complete experiment duration T:",style="Normal")
ins(A_45,"E_PoCol,continuous ≈ Σᵢ Pᵢ(active)·T + Σᵢ Eᵢ(coord).",style="Normal",center=True)
ins(A_45,"Accordingly, deterministic nonce-range separation does not independently imply lower electricity consumption. Energy reduction requires a measurable transition from active hashing to a lower-power state. Continuous operation may improve performance under a tested configuration, but under equal active power and duration it is not expected to consume less energy than a continuously operating PoW baseline.",style="Normal")

# 6.4
for t in [
 "The operating policy must be stated explicitly for each PoCol experiment. A complete comparative design distinguishes among: (1) a continuously operating PoW baseline; (2) PoCol under continuous operation; and (3) PoCol under the post-range idle policy. Energy accounting for the idle scenario must separately record active mining time, idle time, active power, idle power, and coordination overhead.",
 "Parallel PoW, Green-PoW, Proof of Team Sprint, Collaborative Proof-of-Work, StrongChain, and APoW are treated as conceptual and architectural comparisons unless their complete algorithms are independently implemented under identical simulation parameters. Numerical superiority must not be claimed from cross-paper results obtained using different hardware, network sizes, workloads, difficulty settings, or energy models.",
]: ins(A_65,t,style="Normal")

# 7.4.3
ins(A_75,"The energy-saving and continuous-performance policies represent different operating objectives. Under the idle policy, miners stop active hashing after completing their assigned ranges, and energy savings depend on the measured idle power and idle duration. Under the continuous policy, miners remain active and may process further non-overlapping work. If active power and duration are equal to the PoW baseline, this mode should not be interpreted as inherently energy-saving. Any observed throughput advantage must be limited to the implemented configuration and reported together with the corresponding stale-block rate, block interval, active hash rate, and total energy.",style="Normal")

# ---- STEP H: List of Tables entry (red) ----
ins(A_lot,"Table 3.9. Comparison of PoCol with parallel, collaborative, auditable, and energy-aware Proof-of-Work approaches",style="Normal")

# ---- STEP I: mark all fields dirty so Word updates TOC/PAGEREF on open ----
settings = d.settings.element
uf = settings.find(qn('w:updateFields'))
if uf is None:
    uf = OxmlElement('w:updateFields'); settings.append(uf)
uf.set(qn('w:val'),'true')

d.save(FILE)
print("STEP C-I complete. paras now:", len(d.paragraphs), "tables:", len(d.tables))
