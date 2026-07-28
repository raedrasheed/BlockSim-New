#!/usr/bin/env python3
"""Apply verified bibliographic corrections to the thesis bibliography.
Each correction is confirmed against the official publisher record (Crossref/
arXiv/publisher page). Corrected tokens are coloured red (revision marker);
unchanged text stays black. Thesis body text is NOT touched."""
import shutil, copy
import docx
from docx.shared import RGBColor
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

SRC = "Raed-Rasheed-draft-38-Final.docx"
OUT = "Raed-Rasheed-draft-39-References-Verified.docx"
shutil.copy(SRC, OUT)
RED = "FF0000"
d = docx.Document(OUT)
bib = d.tables[-1]

rows = {r.cells[0].text.strip(): r for r in bib.rows}

def apply(label, old, new, redpart):
    """In the single-run bib cell, replace old->new and colour `redpart` red.
    redpart='' means a pure deletion (no red)."""
    cell = rows[label].cells[1]
    p = cell.paragraphs[0]
    # locate the run that contains `old` (entries may already be split by a
    # previous correction to the same entry)
    r = None
    for run in p.runs:
        if old in (run.text or ""):
            r = run; break
    assert r is not None, f"{label}: old not found in any run: {old!r}"
    text = r.text
    idx = text.index(old)
    before = text[:idx]
    after = text[idx+len(old):]
    # base rPr to copy for formatting preservation
    base_rpr = r._r.find(qn('w:rPr'))
    def mk(t, red=False):
        nr = OxmlElement('w:r')
        rpr = copy.deepcopy(base_rpr) if base_rpr is not None else OxmlElement('w:rPr')
        # strip any existing color
        for c in rpr.findall(qn('w:color')): rpr.remove(c)
        if red:
            col = OxmlElement('w:color'); col.set(qn('w:val'), RED); rpr.append(col)
        nr.append(rpr)
        wt = OxmlElement('w:t'); wt.set(qn('xml:space'), 'preserve'); wt.text = t
        nr.append(wt)
        return nr
    # split `new` around redpart
    parent = r._r.getparent()
    r_index = list(parent).index(r._r)
    parent.remove(r._r)
    new_runs = []
    if before: new_runs.append(mk(before, red=False))
    if redpart and redpart in new:
        j = new.index(redpart)
        nb, rp, na = new[:j], redpart, new[j+len(redpart):]
        if nb: new_runs.append(mk(nb, red=False))
        new_runs.append(mk(rp, red=True))
        if na: new_runs.append(mk(na, red=False))
    else:
        # deletion or whole-new: colour whole `new` red only if redpart truthy
        if new: new_runs.append(mk(new, red=bool(redpart)))
    if after: new_runs.append(mk(after, red=False))
    for k, nr in enumerate(new_runs):
        parent.insert(r_index + k, nr)

# ---- verified corrections (label, old, new, red-part) ----
CORR = [
 ("[42]", "A. Manzoor", "A. Mansoor", "Mansoor"),
 ("[87]", "J. L. Y. Tong", "J. L. Y. Terpstra Tong", "Terpstra Tong"),
 ("[87]", "160–177", "136–151", "136–151"),
 ("[93]", "M. P. Asadauskas, C. Cachin, and I. Amores Sesar", "M. P. Asadauskas", ""),
 ("[94]", "M. H. ur Rehman", "M. H. Rehmani", "Rehmani"),
 ("[96]", "2490–2510, 2020", "2490–2510, 2022", "2022"),
 ("[109]", "A. Christian, D. Therry, R.", "A. C. D. Therry, R.", "A. C. D. Therry"),
 ("[112]", "in Applied Cryptography and Network Security Workshops (Lecture Notes in Computer Science), Springer, 2022, doi",
           "in Progress in Cryptology – INDOCRYPT 2021 (Lecture Notes in Computer Science), Springer, 2021, pp. 559–583, doi",
           "Progress in Cryptology – INDOCRYPT 2021 (Lecture Notes in Computer Science), Springer, 2021, pp. 559–583"),
 ("[115]", "(IJACSA) 7 (2024)", "(IJACSA), vol. 15, no. 7, 2024", "vol. 15, no. 7, 2024"),
 ("[117]", "vol. 10, no. 24, 2023", "vol. 11, no. 2, pp. 2855–2869, 2024", "vol. 11, no. 2, pp. 2855–2869, 2024"),
 ("[124]", "no. 7, pp. 517", "no. 7-8, pp. 517", "7-8"),
 ("[125]", "C. Brian", "S. Pirani", "S. Pirani"),
 ("[125]", "700–702, 2020, doi", "700–702, 2018, doi", "2018"),
 ("[126]", "Quasi-Experiments: Resource Flexibility", "Quasi-Experiments: The Resource Flexibility", "The"),
 ("[144]", "F. Raheman,", "F. Raheman and T. Bhagat,", "and T. Bhagat"),
 ("[146]", "O. Zumburidze, N. Adamashvili, R. State, R. Tonelli, and H. Taherdoost", "H. Taherdoost", "H. Taherdoost"),
 ("[147]", "R. M. Ashu and S. Zafar", "A. Gautam, R. Mahajan, and S. Zafar", "A. Gautam, R. Mahajan, and S. Zafar"),
]
for label, old, new, red in CORR:
    apply(label, old, new, red)
    print(f"corrected {label}: {old!r} -> {new!r}")

d.save(OUT)
print("saved", OUT)
