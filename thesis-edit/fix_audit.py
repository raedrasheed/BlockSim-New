#!/usr/bin/env python3
"""Audit correction: restore strict IEEE first-appearance order.
Old ref [133] (Sedlmeir) first appears in Section 3.6.12 (before the inserted
Section 3.7), so it must be [133]; the 5 new refs must be [134]-[138].
Rotation on numbers 133-138 only: {133->134,134->135,135->136,136->137,137->138,138->133}."""
import re
import docx
from docx.oxml.ns import qn
from docx.document import Document as _Doc
from docx.table import _Cell, Table
from docx.text.paragraph import Paragraph

FILE = "Raed-Rasheed-draft-38-Final.docx"
d = docx.Document(FILE)
bib_tbl = d.tables[-1]._tbl
ROT = {133:134, 134:135, 135:136, 136:137, 137:138, 138:133}

def remap(m):
    n = int(m.group(1))
    return f"[{ROT[n]}]" if n in ROT else m.group(0)

def ibi(pr):
    el = pr.element.body if isinstance(pr,_Doc) else (pr._tc if isinstance(pr,_Cell) else pr)
    for ch in el.iterchildren():
        if ch.tag == qn('w:p'): yield Paragraph(ch, pr)
        elif ch.tag == qn('w:tbl'): yield Table(ch, pr)

# ---- 1. in-text citation rotation (paras + tables, excluding bibliography) ----
changed = 0
def fix_par(p):
    global changed
    if 'PAGEREF' in p._p.xml: return
    for r in p.runs:
        if r.text and '[' in r.text:
            new = re.sub(r'\[(\d{1,3})\]', remap, r.text)
            if new != r.text: r.text = new; changed += 1
for b in ibi(d):
    if isinstance(b, Paragraph):
        fix_par(b)
    else:
        if b._tbl is bib_tbl: continue
        for row in b.rows:
            for c in row.cells:
                for p in c.paragraphs: fix_par(p)

# ---- 2. bibliography: relabel the 6 rows + move Sedlmeir row above the new ones ----
bib = d.tables[-1]
rows_by_label = {}
for row in bib.rows:
    lab = row.cells[0].text.strip()
    if lab in ('[132]','[133]','[134]','[135]','[136]','[137]','[138]'):
        rows_by_label[lab] = row

def relabel(row, newlab):
    for r in row.cells[0].paragraphs[0].runs:
        if '[' in r.text:
            r.text = re.sub(r'\[\d{1,3}\]', newlab, r.text, count=1); return
    # if label empty of runs, set first run
    p = row.cells[0].paragraphs[0]
    (p.runs[0] if p.runs else p.add_run('')).text = newlab

sed = rows_by_label['[138]']      # Sedlmeir (existing, black) -> [133]
par = rows_by_label['[133]']      # Parallel PoW -> [134]
# relabel all six (rotation)
relabel(rows_by_label['[133]'], '[134]')  # Parallel
relabel(rows_by_label['[134]'], '[135]')  # StrongChain
relabel(rows_by_label['[135]'], '[136]')  # Collaborative
relabel(rows_by_label['[136]'], '[137]')  # PoTS
relabel(rows_by_label['[137]'], '[138]')  # APoW
relabel(sed, '[133]')                       # Sedlmeir

# move Sedlmeir <w:tr> to just before the Parallel-PoW row (now [134])
sed._tr.getparent().remove(sed._tr)
par._tr.addprevious(sed._tr)

d.save(FILE)
print("in-text citation runs changed:", changed)
print("bibliography rows relabelled/reordered: 6 (Sedlmeir -> [133]; new refs -> [134]-[138])")
