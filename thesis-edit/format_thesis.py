#!/usr/bin/env python3
"""IUG university-format pass (formatting only; no content changes)."""
import re
import docx
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH as AL
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

FILE = "Raed-Rasheed-draft-40-UniversityFormat.docx"
d = docx.Document(FILE)
TNR = "Times New Roman"

def set_font_name(style_or_run_font, name):
    style_or_run_font.name = name

# ---------------------------------------------------------------------------
# 1. Page setup: A4, margins (L 3.5, others 3 cm), no gutter, all sections
# ---------------------------------------------------------------------------
for s in d.sections:
    s.page_width = Cm(21); s.page_height = Cm(29.7)
    s.left_margin = Cm(3.5); s.right_margin = Cm(3.0)
    s.top_margin = Cm(3.0); s.bottom_margin = Cm(3.0)
    try: s.gutter = Cm(0)
    except Exception: pass
    # single-sided: clear mirrorMargins if present
    sectPr = s._sectPr
    for el in sectPr.findall(qn('w:mirrorMargins')):
        sectPr.remove(el)

# ---------------------------------------------------------------------------
# 2. docDefaults -> Times New Roman for latin text (leave cs/eastAsia for Arabic)
# ---------------------------------------------------------------------------
styles_el = d.styles.element
rpr = styles_el.find(qn('w:docDefaults')+'/'+qn('w:rPrDefault')+'/'+qn('w:rPr'))
if rpr is not None:
    rfonts = rpr.find(qn('w:rFonts'))
    if rfonts is None:
        rfonts = OxmlElement('w:rFonts'); rpr.insert(0, rfonts)
    for a in ('w:asciiTheme','w:hAnsiTheme'):
        if rfonts.get(qn(a)) is not None: del rfonts.attrib[qn(a)]
    rfonts.set(qn('w:ascii'), TNR); rfonts.set(qn('w:hAnsi'), TNR)

# ---------------------------------------------------------------------------
# 3. Normal (body): TNR 12 pt, justified, 1.5 line, 6 pt before/after, 1 cm indent
# ---------------------------------------------------------------------------
normal = d.styles['Normal']
normal.font.name = TNR; normal.font.size = Pt(12)
pf = normal.paragraph_format
pf.line_spacing = 1.5
pf.space_before = Pt(6); pf.space_after = Pt(6)
pf.alignment = AL.JUSTIFY
pf.first_line_indent = Cm(1)
# make sure Normal rFonts cs stays (Arabic) - only ascii/hAnsi set by font.name

# ---------------------------------------------------------------------------
# 4. Heading styles
# ---------------------------------------------------------------------------
def style_heading(name, size, align):
    st = d.styles[name]
    st.font.name = TNR; st.font.size = Pt(size); st.font.bold = True
    st.paragraph_format.alignment = align
    # headings: no first-line indent
    st.paragraph_format.first_line_indent = Cm(0)
style_heading('Heading 1', 14, AL.CENTER)   # chapter title
style_heading('Heading 2', 13, AL.LEFT)     # first-level (x.y)
style_heading('Heading 3', 12, AL.LEFT)     # second-level (x.y.z)
style_heading('Heading 4', 12, AL.LEFT)     # deeper (consistent)
# Chapter divider style -> TNR (size left as design element; flagged in report)
cp = d.styles['Chapter Paper']
cp.font.name = TNR
cp.paragraph_format.first_line_indent = Cm(0)

# zero the inherited 1 cm first-line indent on non-body styles that are
# basedOn Normal (captions, dividers, lists, TOC entries)
for st in d.styles:
    nm = st.name or ""
    if nm.startswith('TOC') or nm in ('Caption', 'List Paragraph', 'List Bullet',
                                       'List Number', 'Author', 'key words', 'Subtitle',
                                       'Title'):
        try: st.paragraph_format.first_line_indent = Cm(0)
        except Exception: pass

# ---------------------------------------------------------------------------
# 5. Body captions (Table/Figure in the body): TNR 12 pt, left, no indent
#    (skip the preliminary List of Tables/Figures entries: index < body start)
# ---------------------------------------------------------------------------
body_start = None
for i, p in enumerate(d.paragraphs):
    if p.style.name == 'Chapter Paper':
        body_start = i; break
cap_re = re.compile(r'^\s*(Table|Figure)\s+\d+\.\d+\.\s')  # real captions only (N.N. period)
caps = 0
for i, p in enumerate(d.paragraphs):
    if body_start is not None and i < body_start:
        continue
    if cap_re.match(p.text):
        p.alignment = AL.LEFT
        p.paragraph_format.first_line_indent = Cm(0)
        for r in p.runs:
            r.font.name = TNR
            if r.font.size is None or r.font.size != Pt(12):
                r.font.size = Pt(12)
        caps += 1

# ---------------------------------------------------------------------------
# 6. Remove inherited first-line indent from all table cells (equations, bib,
#    data tables) so only body paragraphs are indented.
# ---------------------------------------------------------------------------
for t in d.tables:
    for row in t.rows:
        for c in row.cells:
            for p in c.paragraphs:
                p.paragraph_format.first_line_indent = Cm(0)

# ---------------------------------------------------------------------------
# 7. Footers: centre the page-number paragraph (bottom-centre numbering)
# ---------------------------------------------------------------------------
for s in d.sections:
    for foot in (s.footer, s.first_page_footer, s.even_page_footer):
        if foot is None: continue
        for p in foot.paragraphs:
            p.alignment = AL.CENTER

# ---------------------------------------------------------------------------
# 8. Ensure Word refreshes TOC / fields on open
# ---------------------------------------------------------------------------
sett = d.settings.element
uf = sett.find(qn('w:updateFields'))
if uf is None:
    uf = OxmlElement('w:updateFields')
    hdr = sett.find(qn('w:hdrShapeDefaults'))
    (hdr.addprevious(uf) if hdr is not None else sett.append(uf))
uf.set(qn('w:val'), 'true')

d.save(FILE)
print("formatting applied. captions formatted:", caps, "| body_start para:", body_start)
