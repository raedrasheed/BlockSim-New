"""Stage 2 of the red-line revision: convert the *added* plain-text mathematics
into native Word (OMML) equations, in place, preserving everything else.

Scope (only the content WE added is plain text; the original equations (1)-(7)
are already OMML and are left untouched):
  * Rebuild the PoW economic prose (sec. D) and PoS prose (sec. E) so every
    symbol is inline OMML, and add displayed, numbered equations (8)-(11) using
    clones of the original equation-table style.
  * Convert the inline "E = P x T" in the Results to OMML (E = P . T).

All added/converted math is coloured RED (C00000). Original text is untouched.

Run AFTER docs/conservative_revision.py. Edits the same redline DOCX in place.
"""

import copy
import os

import docx
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

F = os.path.join(os.path.dirname(__file__),
                 "Extending_BlockSim_Energy_Carbon_REVISED_redline.docx")
RED = "C00000"
MNS = "http://schemas.openxmlformats.org/officeDocument/2006/math"

doc = docx.Document(F)
report = []


# ----------------------------------------------------------------- OMML builder
def M(tag):
    return OxmlElement("m:" + tag)


def mrun(text, up=False, red=True):
    """A single math run <m:r>. up=True -> upright (roman) label."""
    r = M("r")
    rpr = M("rPr")
    sty = M("sty"); sty.set(qn("m:val"), "p" if up else "i"); rpr.append(sty)
    r.append(rpr)
    if red:
        wrpr = OxmlElement("w:rPr")
        col = OxmlElement("w:color"); col.set(qn("w:val"), RED); wrpr.append(col)
        r.append(wrpr)
    t = M("t"); t.set(qn("xml:space"), "preserve"); t.text = text
    r.append(t)
    return [r]


def _box(tag, *childlists):
    el = M(tag)
    for cl in childlists:
        el.extend(cl)
    return el


def sSub(base, sub):
    x = M("sSub"); x.append(_box("e", base)); x.append(_box("sub", sub)); return [x]


def sSup(base, sup):
    x = M("sSup"); x.append(_box("e", base)); x.append(_box("sup", sup)); return [x]


def sSubSup(base, sub, sup):
    x = M("sSubSup")
    x.append(_box("e", base)); x.append(_box("sub", sub)); x.append(_box("sup", sup))
    return [x]


def frac(num, den):
    x = M("f"); x.append(_box("num", num)); x.append(_box("den", den)); return [x]


def nary(chrv, sub, sup, body):
    x = M("nary"); pr = M("naryPr")
    c = M("chr"); c.set(qn("m:val"), chrv); pr.append(c)
    lim = M("limLoc"); lim.set(qn("m:val"), "undOvr"); pr.append(lim)
    grow = M("grow"); grow.set(qn("m:val"), "1"); pr.append(grow)
    x.append(pr)
    x.append(_box("sub", sub)); x.append(_box("sup", sup)); x.append(_box("e", body))
    return [x]


def delim(inner):
    x = M("d"); x.append(_box("e", inner)); return [x]


def omath(children, display=False):
    om = M("oMath"); om.extend(children)
    if display:
        para = M("oMathPara")
        pr = M("oMathParaPr"); jc = M("jc"); jc.set(qn("m:val"), "center"); pr.append(jc)
        para.append(pr); para.append(om)
        return para
    return om


# operators
EQ = lambda: mrun("=", up=True)
PLUS = lambda: mrun("+", up=True)
CDOT = lambda: mrun("·", up=True)   # middle dot
COMMA = lambda: mrun(", ", up=True)

# common symbols (children lists)
SUB_T = lambda: mrun("t")
Bt = lambda: sSub(mrun("B"), SUB_T())
Ft = lambda: sSub(mrun("F"), SUB_T())
Pt = lambda: sSub(mrun("P"), SUB_T())
Rt = lambda: sSub(mrun("R"), SUB_T())
KAPPA = lambda: mrun("κ")           # κ
EPS = lambda: mrun("ε")             # ε
Celec = lambda: sSub(mrun("C"), mrun("elec,", up=True) + mrun("t"))
Ebudget = lambda: sSub(mrun("E"), mrun("budget,", up=True) + mrun("t"))
Etech = lambda: sSub(mrun("E"), mrun("technical", up=True))
Hblock = lambda: sSub(mrun("H"), mrun("block", up=True))
EpsHash = lambda: sSub(mrun("ε"), mrun("hash", up=True))
Dvar = lambda: mrun("D")
Two32 = lambda: sSup(mrun("2", up=True), mrun("32", up=True))
EPoWt = lambda: sSub(mrun("E"), mrun("PoW,", up=True) + mrun("t"))
Eit = lambda: sSub(mrun("E"), mrun("i,t"))
sit = lambda: sSub(mrun("s"), mrun("i,t"))
EPoS = lambda: sSub(mrun("E"), mrun("PoS", up=True))
Pv = lambda: sSub(mrun("P"), mrun("v"))
Tvar = lambda: mrun("T")
uv = lambda: sSub(mrun("u"), mrun("v"))
Vvar = lambda: mrun("V")
Ecomm = lambda: sSub(mrun("E"), mrun("comm", up=True))


# ----------------------------------------------------------- paragraph utilities
def find_para_idx(substr):
    for i, p in enumerate(doc.paragraphs):
        if substr in p.text:
            return i
    return -1


def clear_para_children(p):
    """Remove all runs and inline math, keep pPr."""
    for child in list(p._p):
        if child.tag == qn("w:pPr"):
            continue
        p._p.remove(child)


def red_text_run(text):
    r = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    rf = OxmlElement("w:rFonts")
    rf.set(qn("w:ascii"), "Times New Roman"); rf.set(qn("w:hAnsi"), "Times New Roman")
    rpr.append(rf)
    col = OxmlElement("w:color"); col.set(qn("w:val"), RED); rpr.append(col)
    sz = OxmlElement("w:sz"); sz.set(qn("w:val"), "22"); rpr.append(sz)  # 11pt
    r.append(rpr)
    t = OxmlElement("w:t"); t.set(qn("xml:space"), "preserve"); t.text = text
    r.append(t)
    return r


def render_tokens(p, tokens):
    """tokens: list of ('t', str) or ('m', children-list). Appends to paragraph p."""
    clear_para_children(p)
    for kind, val in tokens:
        if kind == "t":
            p._p.append(red_text_run(val))
        else:
            p._p.append(omath(val, display=False))


def insert_blocks_after(anchor_p, blocks):
    cur = anchor_p._p
    for blk in blocks:
        cur.addnext(blk)
        cur = blk


def empty_para():
    return OxmlElement("w:p")


def eq_table(children, number):
    """Clone original equation table 0 and put `children` as a centered display
    equation in the first cell; set the second cell label to (number), red."""
    new_tbl = copy.deepcopy(doc.tables[0]._tbl)
    tcs = new_tbl.findall(".//" + qn("w:tc"))
    # first cell: replace math
    c0 = tcs[0]
    wp = c0.find(qn("w:p"))
    if wp is None:
        wp = OxmlElement("w:p"); c0.append(wp)
    for op in wp.findall(qn("m:oMathPara")):
        wp.remove(op)
    for om in wp.findall(qn("m:oMath")):
        wp.remove(om)
    wp.append(omath(children, display=True))
    # second cell: number, red
    c1 = tcs[1]
    ts = c1.findall(".//" + qn("w:t"))
    if ts:
        ts[0].text = f"({number})"
        # color its run red
        run = ts[0].getparent()
        rpr = run.find(qn("w:rPr"))
        if rpr is None:
            rpr = OxmlElement("w:rPr"); run.insert(0, rpr)
        for c in rpr.findall(qn("w:color")):
            rpr.remove(c)
        col = OxmlElement("w:color"); col.set(qn("w:val"), RED); rpr.append(col)
        for extra in ts[1:]:
            extra.text = ""
    return new_tbl


# ============================================================ EDIT: PoW (sec. D)
i = find_para_idx("Because PoW energy is bounded by mining economics")
if i < 0:
    raise SystemExit("PoW prose paragraph not found")
pD = doc.paragraphs[i]
tokens_D = [
    ("t", "Because PoW energy is bounded by mining economics, we add an economic "
          "model. The expected per-block reward in fiat is given by Eq. (8), where "),
    ("m", Bt()), ("t", " is the block subsidy, "),
    ("m", Ft()), ("t", " the transaction fees, and "),
    ("m", Pt()), ("t", " the coin price. A rational mining market spends a "
                       "fraction "),
    ("m", KAPPA()), ("t", " ∈ [0,1] of this reward on electricity, giving the "
                          "economic energy budget per block in Eq. (9), where "),
    ("m", Celec()), ("t", " is the electricity price per kWh. When hardware data "
                          "are available, a technical bound "),
    ("m", Etech() + EQ() + Hblock() + CDOT() + EpsHash()),
    ("t", " is derived from the expected hashes per block ("),
    ("m", Hblock() + EQ() + Dvar() + CDOT() + Two32()),
    ("t", ") and the per-hash efficiency "),
    ("m", EpsHash()), ("t", ", and the network per-block energy is "),
    ("m", EPoWt() + EQ() + mrun("min", up=True) + delim(Etech() + COMMA() + Ebudget())),
    ("t", ". The miner-level allocation by hashpower share "),
    ("m", sit()), ("t", " is given by Eq. (10). Thus PoW energy depends on coin "
                        "price, reward, electricity price, efficiency, and "
                        "difficulty/hashrate, and only the distribution depends on "
                        "the number of miners."),
]
render_tokens(pD, tokens_D)
report.append("[OK] rebuilt PoW prose with inline OMML")

# displayed equations (8),(9),(10)
eq8 = Rt() + EQ() + delim(Bt() + PLUS() + Ft()) + CDOT() + Pt()
eq9 = Ebudget() + EQ() + frac(KAPPA() + CDOT() + Rt(), Celec())
eq10 = Eit() + EQ() + sit() + CDOT() + EPoWt()
blocks_D = [eq_table(eq8, 8), empty_para(),
            eq_table(eq9, 9), empty_para(),
            eq_table(eq10, 10), empty_para()]
insert_blocks_after(pD, blocks_D)
report.append("[OK] inserted displayed equations (8),(9),(10)")

# ============================================================ EDIT: PoS (sec. E)
j = find_para_idx("For PoS we add a validator-count model")
if j < 0:
    raise SystemExit("PoS prose paragraph not found")
pE = doc.paragraphs[j]
tokens_E = [
    ("t", "For PoS we add a validator-count model: the network energy is given by "
          "Eq. (11), where "),
    ("m", Pv()), ("t", " is per-validator power (W), "),
    ("m", Tvar()), ("t", " the horizon (h), "),
    ("m", uv()), ("t", " ∈ [0,1] the uptime, "),
    ("m", Vvar()), ("t", " the number of validators, and "),
    ("m", Ecomm()), ("t", " the communication energy. PoS energy therefore scales "
                          "with validator count and hardware power and is "
                          "independent of coin price."),
]
render_tokens(pE, tokens_E)
report.append("[OK] rebuilt PoS prose with inline OMML")

# displayed equation (11): E_PoS = ( sum_{v=1}^{V} (P_v . T . u_v) ) / 1000 + E_comm
nary_body = delim(Pv() + CDOT() + Tvar() + CDOT() + uv())
nary_sum = nary("∑", mrun("v") + mrun("=", up=True) + mrun("1", up=True),
                Vvar(), nary_body)
eq11 = EPoS() + EQ() + frac(nary_sum, mrun("1000", up=True)) + PLUS() + Ecomm()
insert_blocks_after(pE, [eq_table(eq11, 11), empty_para()])
report.append("[OK] inserted displayed equation (11)")

# ===================================================== EDIT: inline E = P . T
k = find_para_idx("fixed-power configurations reproduce")
if k < 0:
    k = find_para_idx("closed-form E = P")
if k >= 0:
    p = doc.paragraphs[k]
    for run in list(p.runs):
        if "E = P" in run.text and "T" in run.text:
            txt = run.text
            old = "E = P × T" if "E = P × T" in txt else (
                  "E = P x T" if "E = P x T" in txt else None)
            if old is None:
                # tolerate any spacing of x/×
                import re
                m = re.search(r"E\s*=\s*P\s*[x×]\s*T", txt)
                old = m.group(0) if m else None
            if old:
                before, after = txt.split(old, 1)
                anchor = run._element
                # order: before(run) | omath | after
                ept = mrun("E") + EQ() + mrun("P") + CDOT() + mrun("T")
                if after:
                    a = red_text_run(after); anchor.addnext(a)
                anchor.addnext(omath(ept, display=False))
                run.text = before
                report.append("[OK] converted inline 'E = P x T' to OMML")
                break
else:
    report.append("[skip] E = P x T anchor not found")

doc.save(F)
print("Saved:", F)
for line in report:
    print("  " + line)
# quick structural check
d2 = docx.Document(F)
print("oMathPara now:", len(d2.element.body.findall('.//' + qn('m:oMathPara'))),
      "(expected 11: 7 original + 4 new)")
