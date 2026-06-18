"""Stage 4 (formatting-only): finalise the red-line revision formatting.

Strictly formatting / placement; no scientific content changes:
  * Replace the 4 EMBEDDED new figures (image7-10) with RED, 10 pt, centered
    placement markers "[Insert Figure N here: ...]"; keep their captions.
  * Export the 4 new figures as separate renamed files in figures_revision/.
  * Set all newly added / revised RED text to 10 pt; justify fully-red body
    paragraphs (not headings, captions, markers, or the reference list).
  * Leave Figures 1-6, all original black text, and all equations untouched.

Run AFTER stages 1-3. Edits the same redline DOCX in place.
"""

import os
import re
import shutil

import docx
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), ".."))
F = os.path.join(ROOT, "docs", "Extending_BlockSim_Energy_Carbon_REVISED_redline.docx")
FIGSRC = os.path.join(ROOT, "results", "figures")
FIGOUT = os.path.join(ROOT, "figures_revision")
RED = RGBColor(0xC0, 0x00, 0x00)

DESC = {7: "PoW energy sensitivity to miner count",
        8: "PoS energy scaling with validator count",
        9: "Computational/consensus energy versus communication energy",
        10: "Carbon-emission sensitivity under different γ values"}
SRC = {7: "fig_pow_energy_vs_miners",
       8: "fig_pos_energy_vs_validators",
       9: "fig_computation_vs_communication",
       10: "fig_carbon_vs_gamma"}
RENAME = {7: "Figure7_PoW_energy_sensitivity",
          8: "Figure8_PoS_validator_scaling",
          9: "Figure9_computation_vs_communication",
          10: "Figure10_carbon_gamma_sensitivity"}

doc = docx.Document(F)
report = []


def is_red(run):
    try:
        return run.font.color is not None and run.font.color.rgb == RED
    except Exception:
        return False


# ----------------------------------------------------------------------------
# A. Export new figures as separate renamed files in figures_revision/
# ----------------------------------------------------------------------------
os.makedirs(FIGOUT, exist_ok=True)
for n in (7, 8, 9, 10):
    for ext in ("png", "pdf"):
        src = os.path.join(FIGSRC, SRC[n] + "." + ext)
        dst = os.path.join(FIGOUT, RENAME[n] + "." + ext)
        shutil.copyfile(src, dst)
report.append(f"[OK] exported 8 files to figures_revision/ ({', '.join(RENAME[n] for n in (7,8,9,10))})")


# ----------------------------------------------------------------------------
# B. Replace embedded NEW figures (image7-10) with red placement markers
# ----------------------------------------------------------------------------
new_img_paras = []
for i, p in enumerate(doc.paragraphs):
    blips = p._p.findall(".//" + qn("a:blip"))
    if not blips:
        continue
    rid = blips[0].get(qn("r:embed"))
    part = doc.part.rels[rid].target_part.partname
    if any(str(part).endswith(f"image{n}.png") for n in (7, 8, 9, 10)):
        new_img_paras.append(i)

for i in new_img_paras:
    p = doc.paragraphs[i]
    cap = doc.paragraphs[i + 1].text if i + 1 < len(doc.paragraphs) else ""
    m = re.match(r"\s*Figure\s+(\d+)\.", cap)
    if not m:
        report.append(f"[WARN] no caption number for image para {i}")
        continue
    n = int(m.group(1))
    # drop the image relationship (orphans the media; not displayed)
    for blip in p._p.findall(".//" + qn("a:blip")):
        try:
            doc.part.drop_rel(blip.get(qn("r:embed")))
        except Exception:
            pass
    # clear runs (removes the drawing) and write the marker
    for r in p._p.findall(qn("w:r")):
        p._p.remove(r)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(f"[Insert Figure {n} here: {DESC[n]}]")
    run.font.name = "Times New Roman"
    run.font.size = Pt(10)
    run.font.color.rgb = RED
    report.append(f"[OK] Figure {n}: embedded image -> placement marker")


# ----------------------------------------------------------------------------
# C. Formatting pass: 10 pt for red runs; justify fully-red body paragraphs
# ----------------------------------------------------------------------------
refs_idx = next((i for i, p in enumerate(doc.paragraphs)
                 if p.text.strip().startswith("References")), len(doc.paragraphs))

justified = 0
sized = 0
for i, p in enumerate(doc.paragraphs):
    style_name = (p.style.name if p.style else "") or ""
    text = p.text.strip()
    is_heading = style_name.startswith("Heading")
    is_caption = bool(re.match(r"Figure\s+(7|8|9|10)\.", text))
    is_marker = text.startswith("[Insert Figure")

    if is_heading:
        # new subsection headings: drop any explicit size override so they
        # inherit the ORIGINAL Heading-2 style exactly (matches old headings).
        for run in p.runs:
            if is_red(run):
                run.font.size = None
    else:
        # body: newly added/revised red text -> 10 pt
        for run in p.runs:
            if is_red(run):
                run.font.size = Pt(10)
                sized += 1

    # justify fully-red body paragraphs (not headings/captions/markers/refs)
    if is_heading or is_caption or is_marker or i >= refs_idx:
        continue
    text_runs = [r for r in p.runs if r.text.strip()]
    if text_runs and all(is_red(r) for r in text_runs):
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        justified += 1

report.append(f"[OK] set 10pt on {sized} red runs; justified {justified} fully-red paragraphs")

doc.save(F)
print("Saved:", F)
for line in report:
    print("  " + line)
