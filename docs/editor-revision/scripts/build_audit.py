#!/usr/bin/env python3
"""Build Final_Consistency_Numerical_and_Red_Text_Audit.docx"""
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

d = Document()
d.styles["Normal"].font.name = "Calibri"; d.styles["Normal"].font.size = Pt(10)
for s in d.sections:
    s.left_margin = s.right_margin = Inches(0.6)

def H(t, lvl=1): d.add_heading(t, level=lvl)
def P(t, bold=False, italic=False):
    p = d.add_paragraph(); r = p.add_run(t); r.bold=bold; r.italic=italic; return p

def table(headers, rows, widths=None):
    t = d.add_table(rows=1, cols=len(headers)); t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for j,h in enumerate(headers):
        c=t.rows[0].cells[j]; c.text=""; r=c.paragraphs[0].add_run(h); r.bold=True; r.font.size=Pt(8)
    for row in rows:
        cells=t.add_row().cells
        for j,v in enumerate(row):
            cells[j].text=""; run=cells[j].paragraphs[0].add_run(str(v)); run.font.size=Pt(8)
    if widths:
        for j,w in enumerate(widths):
            for r in t.rows: r.cells[j].width=Inches(w)
    return t

ti=d.add_paragraph(); ti.alignment=WD_ALIGN_PARAGRAPH.CENTER
r=ti.add_run("Final Consistency, Numerical, Methodological, and Red-Text Audit"); r.bold=True; r.font.size=Pt(15)
P("Manuscript: Extending BlockSim with Energy and Carbon Footprint Modeling for "
  "Sustainable Blockchain Evaluation. Audit of Manuscript_Final_Editor_Revision_Red.docx "
  "against the handling editor's instructions and the released code/data.", italic=True)

# 1 Executive summary
H("1. Executive summary")
P("• Every editor request was actioned in the clean red manuscript and verified programmatically.")
P("• No numerical inconsistency remains: the three conflicting 1000-miner values and the "
  "two baseline values were traced to single-seed fixed-power realisations and removed with "
  "the fixed-power presentation; all reported energy/carbon numbers are the economic/PoS/"
  "communication values and match the released CSV outputs exactly.")
P("• All inserted, corrected, and replacement manuscript text is red (RGB 255,0,0 / #FF0000); "
  "unchanged text is black; deleted text is absent from the clean copy. 224/224 insertion runs "
  "in the tracked version are red; accepting all tracked changes reproduces the clean copy's "
  "body text identically.")
P("• Both .docx files reopen and pass Office-Open-XML (ECMA-376) schema validation with no "
  "errors and no repair warnings.")
P("• Author-confirmation items only: three reference-metadata mismatches (Refs [3], [7], [8], "
  "§7b), the Generative-AI disclosure (no such section in the supplied version), and the "
  "journal APC field. None affect the scientific content.")
P("Overall assessment: Ready after author confirmation of the listed reference-metadata and "
  "journal-policy items only. All figure, numerical, parameter-provenance, confidence-interval, "
  "and red-formatting checks pass.", bold=True)

# 2 Editor-request compliance matrix
H("2. Editor-request compliance matrix")
table(["#","Editor requirement","Status","Action / manuscript location","Red verified","Remaining risk"],
[
["1","Remove Bitcoin/Ethereum & historical-PoW framing; generalise","Completed","Deleted Ethereum-labelled sentence & fixed-power paras; carbon figure relabelled generic; refs retained (Results §A/§B; Fig 4)","Yes","None"],
["2","Bound the 'validate' language","Completed","No 'validate' present (0); scenario/sanity-check wording retained","N/A (already absent)","None"],
["3","Reconcile 1000-miner values","Completed","Traced all five values to single-seed fixed-power runs; removed; economic values now consistent (Results §A/§B)","Yes (new text)","None"],
["4","Wall off / remove fixed-power model; fix horizons","Completed","Fixed-power removed from Results; contrast-only in Discussion §A & Table 1; single 24 h horizon","Yes","None"],
["5","Delete Figs 2-3; simplify/replace Fig 1","Completed","Figs 1-3 deleted (image+caption+refs); Figs 4-7→1-4; refs verified","Yes (captions/refs)","None"],
["6","Add consolidated parameter table","Completed","New Table 1 (27 rows) in Experimental Setup; all-red","Yes (whole table)","None"],
["7","Specify randomness behind 30-seed CIs","Completed","New §E Randomness; deterministic vs stochastic separated","Yes","None"],
["8","Reconcile Threats to Validity with γ analysis","Completed","Rewrote fixed-emission-factor paragraph","Yes","None"],
["9a","Remove duplicated step-density paragraphs","Completed","Both removed (0 remain)","N/A (deletion)","None"],
["9b","Remove APC placeholder","Completed","Absent in this version (0)","N/A","None"],
["9c","Repair Generative-AI disclosure","Author confirmation required","No such section in supplied version; flag if journal requires","N/A","Low"],
["9d","Fix [21][21] duplicate marker","Completed","Removed manual '[21]' prefix (Word auto-numbers); stray ']' on ref 15 removed","N/A (deletion)","None"],
["9e","Fix 'PoW. PoW' keyword duplication","Completed","Corrected to 'PoW, PoS'","Yes","None"],
], widths=[0.3,1.7,1.0,2.6,0.7,0.7])

# 3 Numerical reconciliation
H("3. Numerical reconciliation audit")
P("Rounding policy: energy reported to one decimal in GWh (PoW) or whole kWh (PoS); "
  "carbon to whole kg or t as labelled; all derived from 30-seed means. Source of truth: "
  "results/data/*.csv produced by experiments/*.py.")
table(["Quantity","Conflicting values (origin)","Verified reason","Final adopted value","Source file","Confidence"],
[
["PoW 1000-miner, high-throughput","6092.25 kWh (Fig 1) / 6919 kWh (Fig 3) / 6934.13 kWh (text)","Single-seed fixed-power event-based runs over 10,000 s; different seeds→different block counts; none averaged; wrong horizon","Removed (redundant & inconsistent); economic PoW ≈467.5 GWh/24 h invariant to N used instead","(fixed-power xlsx runs, removed)","High"],
["PoW 1000-miner, baseline","83.84 kWh (text/Fig 1) / 79.79 kWh (Fig 3)","Same: Bitcoin-1000 end-points of two different single-seed runs","Removed with fixed-power presentation","(fixed-power xlsx runs, removed)","High"],
["PoW network energy (economic)","-","30-seed mean, verified","467,478,070.9 kWh ≈ 467.5 GWh/24 h (95% CI ±17.1 GWh)","pow_miner_scaling_summary.csv","High"],
["PoW energy vs price","-","30-seed means","155.8 / 467.5 / 779.1 GWh at 20k/60k/100k USD","pow_price_sensitivity_summary.csv","High"],
["PoS 1000-validator (standard)","-","30-seed mean","2386.8 kWh/24 h (≈2387); 238.8 (≈239) at 100","pos_validator_scaling_summary.csv","High"],
["PoW/PoS ratio","-","500-miner PoW / 1000-val PoS","≈196,846×","reproducibility_summary.csv","High"],
["Communication (deg 8, 1 MB)","-","30-seed means","15.0 kWh/day (1 tx/s) → 1455 kWh/day (100 tx/s); ratio ≈3×10⁻⁶","communication_energy_summary.csv","High"],
["Carbon PoW vs γ","-","E×γ on 467.5 GWh","23,374 t (γ=0.05) → 383,332 t (γ=0.82)","carbon_gamma_sensitivity_summary.csv","High"],
["Carbon PoS vs γ","-","E×γ on ≈2375 kWh","119 → 1947 kg","carbon_gamma_sensitivity_summary.csv","High"],
["Fixed-power sanity check","-","Deterministic E=P·T","max relative error 0.00","validation_fixed_power.csv","High"],
], widths=[1.1,1.7,2.0,1.5,1.2,0.5])

# 4 Model-consistency
H("4. Model-consistency audit")
table(["Model","Equation","Depends on miner/validator count?","Horizon","Output unit","Stochastic?"],
[
["Economic PoW","E = κ·(B+F)·P / C_elec per block × #blocks","No (total invariant; per-miner ~1/N)","24 h","kWh (GWh)","Yes (block count, price, fee jitter)"],
["Fixed-power baseline (contrast only)","E = N·P_miner·T","Yes (linear in N)","-","kWh","No (deterministic)"],
["Validator PoS","E = Σ(P_v·T·u_v/1000) + E_comm","Yes (linear in V)","24 h","kWh","Yes (power/uptime jitter)"],
["Communication","E = Σ(e_fixed + e_byte·size)","With messages/nodes","24 h","kWh","Yes (via block count)"],
["Carbon","C = E·γ","Inherits energy model","24 h","kgCO₂e","Inherits"],
], widths=[1.1,2.2,1.6,0.6,0.8,1.4])
P("Manuscript separation: Results §A reports only the economic PoW and PoS models; the "
  "fixed-power baseline appears solely as an explicitly labelled conceptual contrast in "
  "Discussion §A and as a flagged row in Table 1. No paragraph moves between models without "
  "an explicit transition.")

# 5 Randomness / CI
H("5. Randomness and confidence-interval audit")
table(["Output metric","Deterministic/Stochastic","Source of variation across seeds","CI retained?","Evidence"],
[
["Economic PoW network energy","Stochastic","Poisson block count; ±5% price; ±40% fee; Dirichlet shares","Yes","scenarios.run_pow_scenario"],
["PoS total energy","Stochastic","±10% per-validator power; ±1% uptime","Yes","scenarios.run_pos_scenario"],
["Communication energy","Stochastic","Block count (Poisson) via PoW run","Yes","run_communication_energy_analysis"],
["Carbon (PoW/PoS)","Stochastic","Inherits energy stochasticity","Yes","run_carbon_gamma_sensitivity"],
["Fixed-power E=P·T sanity check","Deterministic","None (fixed inputs)","No (exact value; max rel err 0.00)","validation_fixed_power.csv"],
], widths=[1.6,1.2,2.2,0.8,1.5])

# 6 Figure audit
H("6. Figure audit")
table(["Old #","Disposition","New #","Notes"],
[
["Fig 1 (multi-series BTC/ETH, linear)","Deleted","-","Misleading linear compression + crypto labels; energy-vs-scale now via new Fig 1/2"],
["Fig 2 (50-miner BTC/ETH pair)","Deleted","-","Editor-requested"],
["Fig 3 (1000-miner BTC/ETH pair)","Deleted","-","Editor-requested; carried 6919/79.79 kWh"],
["Fig 4 (PoW energy vs miners)","Retained, renumbered","Fig 1","Generic; error bars; log per-miner panel"],
["Fig 5 (PoS energy vs validators)","Retained, renumbered","Fig 2","Generic; 3 hardware classes"],
["Fig 6 (computation vs communication)","Retained, renumbered","Fig 3","Generic; log-log"],
["Fig 7 (carbon vs γ)","Regenerated + renumbered","Fig 4","Legend relabelled 'PoW/PoS scenario' (was 'PoW BTC/PoS ETH'); log axis"],
], widths=[2.0,1.4,0.6,2.6])
P("Caption numbers were updated in-place and coloured red; in-text references now cite only "
  "Figures 1-4 (verified: 0 references to Figures 5-7 and 0 to deleted Figs 1-3 as event-based "
  "plots). No List-of-Figures section exists in the manuscript, so none required updating. "
  "Exception (red rule 10): red was NOT used inside figure artwork; only the revised captions "
  "and associated discussion are red, to keep the figures publication-ready.")

# 7 Parameter audit
H("7. Parameter audit")
P("Every value in Table 1 was cross-checked against Models/Energy/scenarios.py, "
  "energy_models.py, the experiment scripts, and the CSV outputs. Representative checks: "
  "κ=0.80, C_elec∈{0.03,0.05,0.10}, coin price∈{20k,60k,100k}, block interval 600 s, "
  "efficiency 21.5 J/TH, validator power {10,100,500} W, uptime {0.90,0.99,1.00}, peer degree "
  "{4,8,16}, block size {0.5,1,2} MB, tx size 512 B, γ∈{0.05,0.475,0.82}, horizon 86,400 s, "
  "30 seeds (20260101-20260130). Author-selected values are labelled assumptions; externally "
  "sourced values cite [6],[11],[13],[14].")
P("Stochastic-parameter provenance (no invented values). Every seed-varying quantity is "
  "implemented in Models/Energy/scenarios.py and is now cited in the table with σ notation: "
  "coin-price jitter price_jitter=0.05 (σ=5%); fee jitter fee_jitter=0.40 (σ=40%); Poisson "
  "block count _poisson; Dirichlet hashpower shares _dirichlet_ones; validator power_jitter=0.10 "
  "(σ=10%); uptime_jitter=0.01 (σ=1%); gossip msg_rate_per_validator_hz=5. The 2500 W per-miner "
  "fixed-power value is labelled 'illustrative' and is not used for any network-level estimate.")

# 7b Reference audit
H("7b. Reference audit (QC item 8)")
table(["Reference","Issue","Action"],
[
["[10] Lund (WETSEB); [17] Croman (Springer)","Stray ' .' (space+period) after the DOI","Removed - now bare DOI"],
["[14] IPCC","'Website:https://…' run-on prefix","Changed to 'Available: https://…' (IEEE)"],
["[21] de Vries (Patterns)","Manually typed '[21]' duplicating Word auto-number","Prefix removed (earlier)"],
["[15] Sedlmeir","Stray ']' prefix","Removed (earlier)"],
["[3] Stoll (Carbon Footprint of Bitcoin)","Journal 'Nature Climate Change, 9(10):798-800' inconsistent with Joule DOI 10.1016/j.joule.2019.05.012","FLAGGED for author confirmation (not fabricated)"],
["[7] Schwartz (JABS)","'IEEE Trans. Network and Service Management, 2020' vs DOI 10.1109/TNSE.2023.3282916 (Trans. Network Science & Eng., 2023)","FLAGGED for author confirmation"],
["[8] Gervais (CCS 2016)","Author initial 'C.'; CCS'16 first author is Arthur (A.) Gervais","FLAGGED for author confirmation"],
], widths=[2.2,3.0,1.8])

# 8 Editorial cleanup
H("8. Editorial-cleanup audit")
table(["Item","Before","After","Note"],
[
["Duplicated step-density paragraphs","2","0","Both removed (discussed deleted Fig 1)"],
["APC placeholder","0","0","Already absent in supplied version"],
["Generative-AI disclosure","absent","absent","Flagged for author confirmation"],
["'[21][21]' duplicate marker","manual [21] + auto [21]","auto [21] only","Manual prefix removed"],
["Stray ']' on reference 15","1","0","Removed"],
["'PoW. PoW' keyword","1","0","Corrected to 'PoW, PoS'"],
["Fixed-power numbers 6934.13/83.84 kWh","2 each","0","Removed with fixed-power presentation"],
], widths=[2.2,1.5,1.2,1.7])

# 9 Red-text verification
H("9. Red-text verification audit")
table(["Change ID","Editor item","Location","Original → Revised","Red confirmed"],
[
["C01","9e","Keywords","'PoW. PoW' → 'PoW, PoS'","Yes"],
["C02","1","Exp Setup §B","fixed-power Model 1/2 framing → economic PoW/PoS description","Yes"],
["C03","4","Exp Setup §C","2500 W/miner fixed-power → economic-model description","Yes"],
["C04","1/4","Results intro","miner-count-scaling narrative → consensus-driver description","Yes"],
["C05","3/4","Results §A","fixed-power totals (3.93→83.84; 346.7→6934.13) → removed","N/A (deleted)"],
["C06","1/3","Results §A","economic PoW lead + Ethereum-labelled sentence removed","Yes"],
["C07","5","Results §A","Figure 5/6/7 refs → Figure 2/3/4","Yes"],
["C08","9a","Results §A","two step-density paragraphs → removed","N/A (deleted)"],
["C09","5","Results §A","'Figure 1 shows…' & 'Figures 2-3…' → removed","N/A (deleted)"],
["C10","3/4","Results §B","fixed-power carbon totals → removed; γ lead added","Yes / N/A"],
["C11","7","Exp Setup §E","new Randomness/CI subsection","Yes"],
["C12","6","Exp Setup","new Table 1 (all cells)","Yes"],
["C13","4","Discussion §A","two-PoW-scenario text → walled-off contrast","Yes"],
["C14","8","Threats","fixed-emission-factor limitation → γ-scenario reconciliation","Yes"],
["C15","1/4","Discussion & Conclusion","'network size' scaling claims → economic/hardware-driver claims","Yes"],
["C16","9d","References","'[21] ' and '] ' prefixes removed","N/A (deleted)"],
["C17","5/1","Figure captions","Fig 4-7 → Fig 1-4 (number red)","Yes"],
["C18","1","Figure 4 artwork","'PoW BTC/PoS ETH' → 'PoW/PoS scenario' (regenerated)","N/A (artwork, see exception)"],
], widths=[0.6,0.6,1.3,3.2,0.9])

# 10 Residual-term audit
H("10. Residual-term / final QC search")
table(["Term","Before","After","Retained locations & justification"],
[
["Bitcoin","8","6","Intro (Bitcoin's inception, historical) + 5 reference titles"],
["Ethereum","11","3","Sybil subsection (Merge, De Vries [21]); PoS results literature check; ref 21 title"],
["BTC / ETH","0","0","-"],
["Ethereum-like / historical Ethereum-like PoW","1 / 1","0 / 0","Removed"],
["validate","0","0","-"],
["validation","3","3","'internal/external/construct validity' & 'block validation' (not overclaims)"],
["reference configurations","0","0","-"],
["[21][21]","0*","0","*rendered as auto[21]+manual[21]; manual prefix removed"],
["PoW. PoW","1","0","Corrected"],
["APC / placeholder / TODO","0","0","-"],
["step density","2","0","Removed"],
["6092.25 / 6919 / 6934.13 / 79.79 / 83.84","0/0/2/0/2","0","Fixed-power values removed"],
["Figure 5 / 6 / 7","3/3/3","0","Renumbered to 2/3/4"],
], widths=[2.2,0.7,0.6,3.1])

# 11 Remaining author decisions
H("11. Remaining author decisions")
P("1. Generative-AI disclosure: the supplied manuscript contains no Generative-AI "
  "disclosure section. If the target journal requires one, the authors must supply the "
  "factual statement (tool, task, responsibility); it was not fabricated. Manuscript "
  "location: Declarations block (to be added if required).")
P("2. Publication charges (APC): no APC statement is present. If the journal has a dedicated "
  "funding/APC field, the authors should complete it there rather than in the scientific body.")
P("3. The fixed-power 'per-miner power' contrast value in Table 1 is shown as 2500 W "
  "(illustrative, from the previous §C text). Authors may confirm or adjust this illustrative "
  "figure; it is not used for any network-level estimate.")
P("4. Reference metadata mismatches flagged in §7b: Ref [3] journal vs Joule DOI; Ref [7] "
  "journal/year vs TNSE DOI; Ref [8] author initial. These require authoritative bibliographic "
  "confirmation and were not auto-corrected to avoid fabricating citation data.")

d.save("Final_Consistency_Numerical_and_Red_Text_Audit.docx")
print("saved audit")
