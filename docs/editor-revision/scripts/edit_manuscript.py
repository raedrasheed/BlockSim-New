#!/usr/bin/env python3
"""Stage 1 (clean red): text edits driven by shared_edits.py.
Every inserted/replacement run is RGB(255,0,0)=#FF0000; deletions removed."""
import copy
from docx import Document
from docx.shared import RGBColor
from docx.oxml.ns import qn
import shared_edits as S

SRC = "manuscript.docx"
OUT = "Manuscript_Final_Editor_Revision_Red.docx"
RED = RGBColor(0xFF, 0x00, 0x00)

d = Document(SRC)
paras = d.paragraphs
def wq(t): return qn("w:" + t)

def set_red(run): run.font.color.rgb = RED
def clear_text_runs(p):
    for r in list(p.runs):
        if r._r.find(wq("t")) is not None:
            r._r.getparent().remove(r._r)
def rewrite_red(p, text):
    clear_text_runs(p); run = p.add_run(text); set_red(run)
def delete_paragraph(p): p._p.getparent().remove(p._p)
def _copy_rpr(src, dst):
    s = src._r.find(wq("rPr"))
    if s is not None:
        old = dst._r.find(wq("rPr"))
        if old is not None: dst._r.remove(old)
        dst._r.insert(0, copy.deepcopy(s))
def surgical(p, old, new, red=True):
    for r in list(p.runs):
        if r.text and old in r.text:
            before, after = r.text.split(old, 1)
            src = r; r.text = before
            if new:
                mid = p.add_run(new); _copy_rpr(src, mid)
                if red: set_red(mid)
                aft = p.add_run(after); _copy_rpr(src, aft)
                src._r.addnext(aft._r); src._r.addnext(mid._r)
            else:  # pure deletion
                aft = p.add_run(after); _copy_rpr(src, aft)
                src._r.addnext(aft._r)
            return True
    return False

P = {i: paras[i] for i in range(len(paras))}

# Keywords (para 16): 'PoW. PoW' -> 'PoW, PoS'
kw = P[16].runs
kw[12].text = ","; set_red(kw[12])
kw[14].text = "PoS"; set_red(kw[14])

for idx, text in S.REWRITES.items():
    rewrite_red(P[idx], text)
for idx in S.DELETES:
    delete_paragraph(P[idx])
for idx in S.CLEAR_TEXT:
    clear_text_runs(P[idx])
for idx, old, new in S.SURGICAL:
    assert surgical(P[idx], old, new), f"surgical {idx}"
for idx, prefix in S.PREFIX_DELETE:
    r0 = P[idx].runs[0]
    assert r0.text.startswith(prefix), f"prefix {idx}"
    r0.text = r0.text[len(prefix):]
# reference fixes: delete stray trailing punctuation runs
for idx in S.REF_TRAILING_PUNCT:
    for r in list(P[idx].runs):
        if r.text in (" ", ".", " ."):
            r._r.getparent().remove(r._r)
# reference fixes: replace a whole-run prefix (red)
for idx, old, new in S.REF_REPLACE_RUN:
    done = False
    for r in P[idx].runs:
        if r.text == old:
            r.text = new; set_red(r); done = True; break
    assert done, f"refreplace {idx}"

d.save(OUT)
print("stage-1 saved:", OUT)
