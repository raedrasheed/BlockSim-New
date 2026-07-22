#!/usr/bin/env python3
"""Build FINAL_PRE_SUBMISSION_QC_REPORT.docx"""
from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

d = Document()
d.styles["Normal"].font.name = "Calibri"; d.styles["Normal"].font.size = Pt(10)
for s in d.sections:
    s.left_margin = s.right_margin = Inches(0.7)

def H(t, l=1): d.add_heading(t, level=l)
def P(t, b=False):
    p = d.add_paragraph(); r = p.add_run(t); r.bold = b; return p
def table(headers, rows, widths=None):
    t = d.add_table(rows=1, cols=len(headers)); t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for j, h in enumerate(headers):
        c = t.rows[0].cells[j]; c.text = ""; r = c.paragraphs[0].add_run(h); r.bold = True; r.font.size = Pt(8)
    for row in rows:
        cs = t.add_row().cells
        for j, v in enumerate(row):
            cs[j].text = ""; rn = cs[j].paragraphs[0].add_run(str(v)); rn.font.size = Pt(8)
    if widths:
        for j, w in enumerate(widths):
            for r in t.rows: r.cells[j].width = Inches(w)

ti = d.add_paragraph(); ti.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = ti.add_run("FINAL PRE-SUBMISSION QC REPORT"); r.bold = True; r.font.size = Pt(16)
P("Manuscript: Extending BlockSim with Energy and Carbon Footprint Modeling for "
  "Sustainable Blockchain Evaluation. Final journal-acceptance QC pass on the "
  "already-revised manuscript (targeted corrections only; no rewrite).")

H("1. Issues corrected")
table(["#", "Item", "What was done", "Verified"],
[
["1", "Figures 2 and 3", "The editor's ORIGINAL misleading Figures 1-3 (Bitcoin/Ethereum, linear-compressed overlays carrying the 6919/6092 kWh conflicts) were already removed and the survivors renumbered. Confirmed with the author to KEEP the current Figures 1-4 (PoW-vs-miners, PoS-vs-validators, computation-vs-communication, carbon-vs-grid) - these are the reviewer-requested figures. Cross-references and captions verified: only Figures 1-4 are referenced; no reference to any deleted figure remains.", "Yes"],
["2", "Figure 1", "The old misleading Figure 1 was replaced by the current Figure 1 (PoW energy vs miners; economic model; error bars; log per-miner panel) - a clear, non-compressed figure. No further change needed per author confirmation.", "Yes"],
["3", "Invented parameters", "Every stochastic parameter in Table 1 was traced to Models/Energy/scenarios.py: coin-price jitter (price_jitter=0.05), fee jitter (fee_jitter=0.40), Poisson block count (_poisson), Dirichlet shares (_dirichlet_ones), validator power jitter (power_jitter=0.10), uptime jitter (uptime_jitter=0.01), gossip rate (msg_rate=5). Table wording changed to Gaussian sigma notation with the exact source file/symbol for each row; the 2500 W per-miner value is labelled 'illustrative'. No invented parameters remain.", "Yes"],
["4", "Numerical values", "Re-audited against results/data/*.csv: 467.5 GWh (467,478,070.9), +/-17.1 GWh (17,132,817.6), 155.8/779.1 GWh (price sweep), 2387 kWh (2386.8), 239 kWh (238.8), 196,846x (reproducibility_summary), E=P.T max rel err 0.00. All match to the stated precision.", "Yes"],
["5", "Confidence intervals", "The only CI attached to a value is '(95% CI +/-17.1 GWh)' on PoW network energy - a stochastic quantity (price/fee/block-count jitter). Deterministic E=P.T is reported exactly with no CI. No meaningless CIs remain.", "Yes"],
["6", "AI-style writing", "Shortened and de-duplicated the newly added paragraphs (Randomness subsection; economic-PoW continuation; Discussion). Removed repeated 'invariant/1-over-N' statements; preserved scientific meaning and the authors' voice.", "Yes"],
["7", "Experimental Setup length", "Sections B, C and E condensed; Table 1 retained. Explanatory prose reduced without losing any parameter.", "Yes"],
["8a", "References - mechanical", "Removed stray trailing ' .' after two DOIs (WETSEB 2019; Springer 978-3-662-53357-4_8); changed 'Website:https://...' to 'Available: https://...' for the IPCC citation; earlier removed the duplicated '[21]' prefix and a stray ']' prefix.", "Yes"],
])

H("2. Remaining issues / values requiring author confirmation")
table(["Item", "Issue", "Why not auto-fixed"],
[
["Ref [3] Stoll et al.", "Journal cited as 'Nature Climate Change, vol. 9, no. 10, pp. 798-800' but the DOI 10.1016/j.joule.2019.05.012 is a Joule DOI.", "Correcting journal/volume/pages requires authoritative bibliographic data; not fabricated."],
["Ref [7] Schwartz et al. (JABS)", "Cited as 'IEEE Trans. Network and Service Management, vol. 17, 2020' but DOI 10.1109/TNSE.2023.3282916 is IEEE Trans. Network Science and Engineering (2023).", "Journal/year mismatch; author must confirm the intended source."],
["Ref [8] Gervais et al.", "Author initial given as 'C. Gervais'; the CCS 2016 paper's first author is Arthur Gervais.", "Initial correction left to author to avoid altering a citation without confirmation."],
["Generative-AI disclosure", "No such section exists in the supplied manuscript.", "Add the factual statement only if the target journal requires one; must not be fabricated."],
["APC / publication charges", "No APC statement in the body (previously-removed placeholder).", "Complete the journal's dedicated funding/APC field if applicable."],
["Fixed-power contrast value", "Table 1 lists the per-miner fixed-power baseline as 2500 W (illustrative).", "Author may confirm/adjust; not used for any network-level estimate."],
])

H("3. Response-letter consistency (editor item)")
P("The current response deliverable (Response_to_Handling_Editor.docx) matches "
  "the final manuscript exactly on values, figure numbering (Figures 1-4), and "
  "actions. The earlier per-reviewer response letters (previous round) describe a "
  "superseded figure numbering (old Fig 1-6 including a 'PoW vs PoS orders' "
  "figure) and the mutually inconsistent 'Figures 1-3 kept unchanged' vs "
  "'crypto framing removed' commitments; these are explicitly reconciled and "
  "superseded by the editor-mediated response and the final figure set, as the "
  "handling editor instructed. No statement in the current response letter "
  "conflicts with the clean manuscript.")

H("4. Formatting integrity")
P("- All inserted/corrected/replacement manuscript text is red (RGB 255,0,0 / #FF0000); unchanged text is black; deleted text absent from the clean copy.")
P("- Consolidated Table 1: 196/196 text runs red (fully new table). Figure captions 1-4: figure-number runs red.")
P("- Tracked version: accepting all changes reproduces the clean copy's body text IDENTICALLY; all insertion runs are red.")
P("- Both .docx pass ECMA-376 (Office Open XML) schema validation with no errors and reopen without repair warnings.")

H("5. Final recommendation")
P("Ready after author confirmation of listed items only.", b=True)
P("All figure, numerical, parameter-provenance, confidence-interval, and red-formatting checks pass and the manuscript is internally consistent. The only open items are the three reference-metadata mismatches (Refs [3], [7], [8]) and the journal-policy items (Generative-AI disclosure, APC field), none of which can be resolved without author/journal input and none of which affect the scientific content.")

d.save("FINAL_PRE_SUBMISSION_QC_REPORT.docx")
print("saved QC report")
