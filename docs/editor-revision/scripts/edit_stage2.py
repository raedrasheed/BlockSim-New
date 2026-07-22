#!/usr/bin/env python3
"""Stage 2: figure removal, caption renumbering, parameter table, randomness
subsection. Operates on the stage-1 output in place."""
import copy
from docx import Document
from docx.shared import RGBColor, Pt
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

FILE = "Manuscript_Final_Editor_Revision_Red.docx"
RED = RGBColor(0xFF, 0x00, 0x00)
R = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"

d = Document(FILE)
root = d.element
def wq(t): return qn("w:" + t)

# ---------------------------------------------------------------------------
# A. Remove old Figures 1-3 : image drawings rId8/9/10 + caption AlternateContents
# ---------------------------------------------------------------------------
DEL_EMBEDS = {"rId8", "rId9", "rId10"}
DEL_CAP_PREFIX = ("Figure 1:", "Figure 2:", "Figure 3:")

removed_imgs, removed_caps = [], []

# remove caption AlternateContents
for ac in list(root.iter("{http://schemas.openxmlformats.org/markup-compatibility/2006}AlternateContent")):
    cap = "".join(t.text or "" for t in ac.iter(wq("t")))
    if cap.startswith(DEL_CAP_PREFIX):
        run = ac.getparent()          # the w:r
        run.getparent().remove(run)   # remove whole run
        removed_caps.append(cap[:30])

# remove image drawings (the run holding the w:drawing with the target blip)
for dr in list(root.iter(wq("drawing"))):
    blip = dr.find(".//" + qn("a:blip"))
    emb = blip.get(R + "embed") if blip is not None else None
    if emb in DEL_EMBEDS:
        run = dr.getparent()          # w:r
        run.getparent().remove(run)
        removed_imgs.append(emb)

print("removed caption boxes:", removed_caps)
print("removed image runs   :", removed_imgs)

# ---------------------------------------------------------------------------
# B. Renumber kept captions 4->1, 5->2, 6->3, 7->4 ; colour the number red
# ---------------------------------------------------------------------------
CAP_MAP = {"Figure 4.": "Figure 1.", "Figure 5.": "Figure 2.",
           "Figure 6.": "Figure 3.", "Figure 7.": "Figure 4."}

def redden_run(r_el):
    rpr = r_el.find(wq("rPr"))
    if rpr is None:
        rpr = OxmlElement("w:rPr"); r_el.insert(0, rpr)
    for c in rpr.findall(wq("color")):
        rpr.remove(c)
    col = OxmlElement("w:color"); col.set(qn("w:val"), "FF0000"); rpr.append(col)

caps_done = []
for ac in list(root.iter("{http://schemas.openxmlformats.org/markup-compatibility/2006}AlternateContent")):
    full = "".join(t.text or "" for t in ac.iter(wq("t")))
    for old, new in CAP_MAP.items():
        if full.startswith(old):
            # each w:t holds the whole caption; split its run into red number + black rest
            for wt in list(ac.iter(wq("t"))):
                if wt.text and wt.text.startswith(old):
                    rest = wt.text[len(old):]
                    r_el = wt.getparent()            # w:r
                    # set this run to the red number
                    wt.text = new
                    redden_run(r_el)
                    # create a sibling run (black) with the remaining text
                    new_r = copy.deepcopy(r_el)
                    # remove color from the copy
                    rpr = new_r.find(wq("rPr"))
                    if rpr is not None:
                        for c in rpr.findall(wq("color")):
                            rpr.remove(c)
                    new_r.find(wq("t")).text = rest
                    new_r.find(wq("t")).set(qn("xml:space"), "preserve")
                    r_el.addnext(new_r)
            caps_done.append(f"{old}->{new}")
            break
print("captions renumbered:", caps_done)

# ---------------------------------------------------------------------------
# C. Insert consolidated parameter table (Table 1) after Exp-Setup intro
# ---------------------------------------------------------------------------
def find_para(prefix):
    for p in d.paragraphs:
        if p.text.strip().startswith(prefix):
            return p
    raise RuntimeError("not found: " + prefix)

import shared_edits as S
intro = find_para("To evaluate the effectiveness of the proposed energy and carbon")

HEADER = S.TABLE_HEADER
ROWS = S.TABLE_ROWS

tbl = d.add_table(rows=1 + len(ROWS), cols=len(HEADER))
try:
    tbl.style = d.styles["Table Grid"]
except KeyError:
    pass

def set_cell(cell, text, bold=False):
    cell.text = ""
    p = cell.paragraphs[0]
    run = p.add_run(text)
    run.font.color.rgb = RED           # whole table is NEW -> red
    run.font.size = Pt(8)
    if bold:
        run.font.bold = True

for j, h in enumerate(HEADER):
    set_cell(tbl.rows[0].cells[j], h, bold=True)
for i, row in enumerate(ROWS, start=1):
    for j, val in enumerate(row):
        set_cell(tbl.rows[i].cells[j], val)

# lead sentence + caption (red, new)
lead = d.add_paragraph()
lead.style = d.styles["RSR"]
lr = lead.add_run(S.PARAM_LEAD)
lr.font.color.rgb = RED

cap = d.add_paragraph()
cap.style = d.styles["RSR"]
cr = cap.add_run(S.PARAM_CAPTION)
cr.font.color.rgb = RED
cr.font.bold = True

spacer = d.add_paragraph()

# move lead, table, caption, spacer to right after the intro paragraph
intro._p.addnext(lead._p)
lead._p.addnext(tbl._tbl)
tbl._tbl.addnext(cap._p)
cap._p.addnext(spacer._p)

# ---------------------------------------------------------------------------
# D. Insert 'E. Randomness, Replications, and CI Construction' subsection
#    after the §D evaluation-metrics paragraph
# ---------------------------------------------------------------------------
metrics = find_para("The subsequent metrics are gathered for each simulation")

head = d.add_paragraph()
head.style = d.styles["Heading 2"]
hr = head.add_run(S.RAND_HEAD)
hr.font.color.rgb = RED

body = d.add_paragraph()
body.style = d.styles["RSR"]
br = body.add_run(S.RAND_BODY)
br.font.color.rgb = RED

metrics._p.addnext(head._p)
head._p.addnext(body._p)

d.save(FILE)
print("stage-2 saved:", FILE)
print("paras now:", len(d.paragraphs), "tables now:", len(d.tables))
