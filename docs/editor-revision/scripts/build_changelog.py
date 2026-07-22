#!/usr/bin/env python3
"""Build Revision_Change_Log.xlsx"""
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side

wb = Workbook(); ws = wb.active; ws.title = "Change Log"
hdr = ["Change ID","Editor comment #","File","Section","Paragraph / location",
       "Original text or value","Revised text or value","Reason","Evidence",
       "Red formatting confirmed","Confidence"]
ws.append(hdr)

rows = [
 ["C01","9e","Manuscript (clean+tracked)","Keywords","Keywords line",
  "PoW. PoW","PoW, PoS","Duplicate keyword + wrong separator","Editor letter; keyword line","Yes","High"],
 ["C02","1,4","Manuscript","Experimental Setup §B","Consensus & Workload Parameters",
  "Two PoW timing scenarios (Model 1/Model 2), miner population varied, ~10 min / ~12-13 s",
  "Economic PoW (600 s) + validator PoS; miner {50..1000}, validator {100..5000}; 24 h; refers to Table 1",
  "Remove fixed-power scenario framing; single horizon","scenarios.py; run_*_scaling.py","Yes","High"],
 ["C03","4","Manuscript","Experimental Setup §C","Energy & Carbon Parameters",
  "PoW baseline hashrate/efficiency vs high-throughput fixed 2500 W/miner",
  "Economic PoW (Eqs 5-7) + PoS (Eq 8); γ scenarios",
  "Remove fixed-power derivation","scenarios.py","Yes","High"],
 ["C04","1,4","Manuscript","Results (intro)","Results & Analysis opening",
  "Across all tested network sizes, total energy and CO2 increased with network scale; two-PoW-scenario magnitude sentence",
  "Consensus-aware drivers described; magnitude sentence deleted",
  "Contradicted economic-model invariance","pow_miner_scaling_summary.csv","Yes / N/A","High"],
 ["C05","3,4","Manuscript","Results §A","fixed-power totals bullets",
  "PoW baseline 3.93→83.84 kWh; high-throughput 346.70→6934.13 kWh (10,000 s)",
  "(removed)","Single-seed fixed-power realisations; inconsistent; redundant","xlsx runs; editor letter","N/A (deleted)","High"],
 ["C06","1,3","Manuscript","Results §A","economic PoW paragraphs",
  "'…Ethereum-labelled curve is a historical Ethereum-like PoW configuration…'; ~83-88× fixed-power text",
  "Economic PoW: 467.5 GWh/24 h invariant (±17.1 GWh); 155.8→779.1 GWh vs price; Figure 1",
  "Remove crypto framing; report economic model","pow_miner_scaling/price CSVs","Yes","High"],
 ["C07","5","Manuscript","Results §A/§B","figure references",
  "Figure 5 / Figure 6 / Figure 7","Figure 2 / Figure 3 / Figure 4",
  "Renumber after deleting Figs 1-3","figure inventory","Yes","High"],
 ["C08","9a","Manuscript","Results §A","step-density paragraphs",
  "Two near-identical 'step density' paragraphs (one BTC/ETH, one generic)","(removed)",
  "Duplicated; discussed deleted event-based Fig 1","editor letter","N/A (deleted)","High"],
 ["C09","5","Manuscript","Results §A","figure-1/2-3 discussion",
  "'Figure 1. shows the event-based cumulative energy…'; 'Each of Figures 2-3 compares…'","(removed)",
  "Figures deleted","editor letter","N/A (deleted)","High"],
 ["C10","3,4","Manuscript","Results §B","carbon paragraphs",
  "Fixed-power carbon 1.75→37.31 & 154.28→3085.69 kg; constant-grid scaling text",
  "Carbon lead (C=E·γ, Eq 9); γ range 23,374 t→383,332 t (PoW), 119-1947 kg (PoS); Figure 4",
  "Remove fixed-power carbon; report γ sensitivity","carbon_gamma_sensitivity_summary.csv","Yes / N/A","High"],
 ["C11","7","Manuscript","Experimental Setup §E (new)","Randomness, Replications & CI",
  "(none)","New subsection: deterministic equations; Poisson block count; ±5% price, ±40% fee, Dirichlet shares (PoW); ±10% power, ±1% uptime (PoS); Student-t 95% CI",
  "State source of 30-seed CIs","scenarios.py","Yes","High"],
 ["C12","6","Manuscript","Experimental Setup (new Table 1)","Table 1",
  "(parameters scattered in prose)","Consolidated 27-row parameter table (all-red)",
  "Consolidate parameters for replication","scenarios.py; experiment scripts; CSVs","Yes","High"],
 ["C13","4","Manuscript","Discussion §A","two-PoW-scenario paragraph",
  "high-throughput fixed high per-miner power … larger than baseline …",
  "Naïve fixed-power baseline walled off; economic model reported; contrast only",
  "Prevent model confusion","editor letter","Yes","High"],
 ["C14","8","Manuscript","Threats to Validity","fixed-emission-factor paragraph",
  "The use of a fixed electricity emission factor … results reflect average conditions",
  "Three discrete γ scenarios; not full temporal/geographic variability; scenario-based comparisons",
  "Reconcile with γ sensitivity","carbon_gamma_sensitivity","Yes","High"],
 ["C15","1,4","Manuscript","Discussion & Conclusion","scaling claims",
  "predictable scaling in relation to network size…; scale predictably with network size…",
  "…economic and hardware drivers of each consensus mechanism…",
  "Consistency with invariance result","pow_miner_scaling","Yes","High"],
 ["C16","9d","Manuscript","References","ref 21 and ref 15",
  "'[21] A. de Vries…'; '] M. Sedlmeir…'","'A. de Vries…'; 'M. Sedlmeir…'",
  "Remove manual duplicate number & stray bracket (Word auto-numbers)","reference list","N/A (deleted)","High"],
 ["C17","5,1","Manuscript","Figure captions","captions of kept figures",
  "Figure 4./5./6./7.","Figure 1./2./3./4. (number red)",
  "Renumber captions","figure inventory","Yes","High"],
 ["C18","1","Figure artwork + script","Figure 4 (carbon)","results/figures/fig_carbon_vs_gamma_generic.png",
  "legend 'PoW BTC' / 'PoS ETH'","legend 'PoW scenario' / 'PoS scenario'",
  "Remove crypto labels from figure","regenerate_carbon_figure_generic.py; carbon CSV","N/A (artwork)","High"],
 ["C19","QC-3","Manuscript","Experimental Setup (Table 1)","Table 1 stochastic rows",
  "±5% / ±40% / ±10% / ±1% 'across seeds'; generic 'assumption' sources",
  "Gaussian σ=5/40/10/1% with exact source symbols (price_jitter, fee_jitter, power_jitter, uptime_jitter); added Poisson block-count and Dirichlet-shares rows; each row cites scenarios.py / experiment script",
  "Make every parameter traceable; no invented values","Models/Energy/scenarios.py","Yes","High"],
 ["C20","QC-6,7","Manuscript","Experimental Setup §E; Results §A; Discussion","Randomness subsection; economic-PoW continuation; Discussion",
  "verbose AI-style paragraphs with repetition","condensed prose; removed duplicated 'invariant/1-over-N' statements",
  "Shorten AI-style writing; reduce repetition","-","Yes","High"],
 ["C21","QC-7","Manuscript","Experimental Setup §B, §C","Consensus/Workload; Energy/Carbon params",
  "longer explanatory paragraphs","condensed; Table 1 retained",
  "Shorten Experimental Setup","-","Yes","High"],
 ["C22","QC-8","Manuscript","References","refs [10], [17] DOI lines",
  "'…00009 .' and '…53357-4_8 .' (stray trailing space+period)","'…00009' and '…53357-4_8'",
  "Remove duplicated trailing punctuation","reference list","N/A (deletion)","High"],
 ["C23","QC-8","Manuscript","References","ref [14] IPCC URL line",
  "'Website:https://www.ipcc.ch/…'","'Available: https://www.ipcc.ch/…'",
  "IEEE-consistent URL prefix; fix run-on","reference list","Yes","High"],
]
for r in rows: ws.append(r)

# formatting
head_fill = PatternFill("solid", fgColor="1F4E79")
thin = Side(style="thin", color="BBBBBB"); border = Border(thin,thin,thin,thin)
for c in ws[1]:
    c.font = Font(bold=True, color="FFFFFF", size=9); c.fill = head_fill
    c.alignment = Alignment(wrap_text=True, vertical="top", horizontal="center")
widths = [9,11,16,20,20,34,40,26,24,14,10]
from openpyxl.utils import get_column_letter
for i,w in enumerate(widths,1): ws.column_dimensions[get_column_letter(i)].width = w
for row in ws.iter_rows(min_row=2):
    for c in row:
        c.alignment = Alignment(wrap_text=True, vertical="top"); c.font = Font(size=9); c.border = border
ws.freeze_panes = "A2"

# second sheet: numerical reconciliation
ws2 = wb.create_sheet("Numerical Reconciliation")
ws2.append(["Quantity","Conflicting values","Reason","Final adopted value","Source file","Confidence"])
nrows = [
 ["PoW 1000-miner high-throughput","6092.25 / 6919 / 6934.13 kWh","single-seed fixed-power runs, 10,000 s, none averaged","removed; economic 467.5 GWh/24 h used","fixed-power xlsx (removed)","High"],
 ["PoW 1000-miner baseline","83.84 / 79.79 kWh","two different single-seed runs","removed","fixed-power xlsx (removed)","High"],
 ["PoW network energy (economic)","-","30-seed mean","467,478,070.9 kWh ≈467.5 GWh (±17.1 GWh)","pow_miner_scaling_summary.csv","High"],
 ["PoW energy vs price","-","30-seed means","155.8/467.5/779.1 GWh (20k/60k/100k USD)","pow_price_sensitivity_summary.csv","High"],
 ["PoS 1000-validator standard","-","30-seed mean","2386.8 kWh (≈2387); 238.8 (≈239) at 100","pos_validator_scaling_summary.csv","High"],
 ["PoW/PoS ratio","-","500-miner / 1000-val","196,846×","reproducibility_summary.csv","High"],
 ["Communication deg8 1MB","-","30-seed means","15.0→1455 kWh/day; ratio 3e-06","communication_energy_summary.csv","High"],
 ["Carbon PoW vs γ","-","E×γ","23,374 t→383,332 t","carbon_gamma_sensitivity_summary.csv","High"],
 ["Carbon PoS vs γ","-","E×γ","119→1947 kg","carbon_gamma_sensitivity_summary.csv","High"],
 ["E=P·T sanity check","-","deterministic","max rel error 0.00","validation_fixed_power.csv","High"],
]
for r in nrows: ws2.append(r)
for c in ws2[1]:
    c.font = Font(bold=True, color="FFFFFF", size=9); c.fill = head_fill
    c.alignment = Alignment(wrap_text=True, vertical="top", horizontal="center")
for i,w in enumerate([26,22,30,30,30,10],1): ws2.column_dimensions[get_column_letter(i)].width=w
for row in ws2.iter_rows(min_row=2):
    for c in row: c.alignment=Alignment(wrap_text=True,vertical="top"); c.font=Font(size=9); c.border=border
ws2.freeze_panes="A2"

wb.save("Revision_Change_Log.xlsx")
print("saved change log")
