#!/usr/bin/env python3
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
d=Document(); d.styles["Normal"].font.name="Calibri"; d.styles["Normal"].font.size=Pt(11)
t=d.add_paragraph(); t.alignment=WD_ALIGN_PARAGRAPH.CENTER
r=t.add_run("Final Acceptance-Readiness Checklist"); r.bold=True; r.font.size=Pt(15)
d.add_paragraph("Manuscript: Extending BlockSim with Energy and Carbon Footprint "
                "Modeling for Sustainable Blockchain Evaluation").italic=True
items=[
("All editor comments addressed","Yes"),
("No new experiment introduced","Yes"),
("No unsupported claim added","Yes"),
("Bitcoin/Ethereum framing removed where required","Yes (only justified literature/historical uses retained)"),
("BTC and ETH labels removed","Yes (0 occurrences)"),
("Validation wording bounded","Yes ('validate' absent; scenario/sanity-check language)"),
("1000-miner numerical values consistent everywhere","Yes (conflicting fixed-power values removed)"),
("Different time horizons reconciled/normalised","Yes (single 24 h horizon)"),
("Economic PoW and fixed-power models clearly separated","Yes (fixed-power = contrast only)"),
("Figures 2 and 3 deleted","Yes"),
("Figure 1 removed/simplified/replaced","Yes (removed; carried by new Fig 1/2)"),
("Later figures renumbered","Yes (4-7 → 1-4)"),
("All cross-references valid","Yes (refs cite only Figures 1-4)"),
("Consolidated parameter table added","Yes (Table 1, all-red)"),
("Randomness across seeds documented","Yes (new §E)"),
("Deterministic quantities show no artificial CIs","Yes (E=P·T reported exact)"),
("Threats to Validity matches γ-sensitivity analysis","Yes"),
("Duplicate step-density text removed","Yes"),
("APC placeholder removed","Yes (already absent)"),
("Generative-AI disclosure repaired or flagged","Flagged for author confirmation (no such section present)"),
("'[21][21]' corrected","Yes (manual prefix removed)"),
("'PoW. PoW' corrected","Yes ('PoW, PoS')"),
("No placeholder/truncated text remains","Yes (0 TODO/placeholder)"),
("All modified manuscript text is red","Yes"),
("Unchanged text remains black","Yes"),
("Deleted text absent from clean manuscript","Yes"),
("Red colour is RGB 255,0,0 / #FF0000","Yes"),
("Files reopen without errors (schema-valid)","Yes (ECMA-376 validation passed)"),
("Response report matches the clean manuscript","Yes"),
("Audit report complete","Yes"),
("Change log complete","Yes (18 changes + numerical reconciliation)"),
]
tb=d.add_table(rows=1,cols=2); tb.style="Table Grid"
h=tb.rows[0].cells; h[0].text=""; h[0].paragraphs[0].add_run("Checklist item").bold=True
h[1].text=""; h[1].paragraphs[0].add_run("Status").bold=True
for it,st in items:
    c=tb.add_row().cells; c[0].text=it; c[1].text=st
d.add_paragraph()
fp=d.add_paragraph(); fr=fp.add_run("Final assessment: Ready after minor author confirmation "
    "(Generative-AI disclosure / journal APC field only)."); fr.bold=True
d.save("Acceptance_Readiness_Checklist.docx"); print("saved checklist")
