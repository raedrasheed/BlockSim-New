"""Stage 5 (numbering/cleanup only): correct equation numbering, delete old
uncited Figures 4-6, and renumber the new figures.

No scientific content rewrite. Preserves formatting, red text, sizes, alignment,
captions style and OMML.

Equation renumber (by order of appearance):
    (8)->(5) (9)->(6) (10)->(7) (11)->(8)   [added PoW/PoS eqs]
    (5)->(9) (6)->(10) (7)->(11)            [original carbon eqs shift down]
    (1)-(4) unchanged
Prose refs Eq.(8/9/10/11) updated to (5/6/7/8).

Figures:
    Delete old Figure 4 (image4), Figure 5 (image5), Figure 6 (image6) and the
    two orphaned reference sentences. Renumber new figures 7->4, 8->5, 9->6,
    10->7 (markers, captions, in-text references). Rename the separate files.

Run AFTER stages 1-4. Edits the same DOCX in place.
"""

import os
import re
import docx
from docx.oxml.ns import qn

ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), ".."))
F = os.path.join(ROOT, "docs", "Extending_BlockSim_Energy_Carbon_REVISED_redline.docx")
FIGREV = os.path.join(ROOT, "figures_revision")

doc = docx.Document(F)
report = []

# ---------------------------------------------------------------- 1. EQUATIONS
EQ_MAP = {5: 9, 6: 10, 7: 11, 8: 5, 9: 6, 10: 7, 11: 8}  # 1-4 unchanged
eq_changes = []
for t in doc.tables:
    om = t._tbl.findall(".//" + qn("m:oMath"))
    if not om:
        continue
    cell = t.rows[0].cells[1]
    m = re.search(r"\((\d+)\)", cell.text)
    if not m:
        continue
    old = int(m.group(1))
    new = EQ_MAP.get(old, old)
    if new != old:
        # set the number run text, preserve formatting/colour
        for ts in cell._element.findall(".//" + qn("w:t")):
            if ts.text and f"({old})" in ts.text:
                ts.text = ts.text.replace(f"({old})", f"({new})")
                break
    eq_changes.append((old, new))
report.append(f"[OK] equation labels remapped: {eq_changes}")

# prose equation references: Eq. (8/9/10/11) -> (5/6/7/8)
PROSE_EQ = {8: 5, 9: 6, 10: 7, 11: 8}
def _eq_ref_sub(s):
    return re.sub(r"Eq\.\s*\((\d+)\)",
                  lambda mm: f"Eq. ({PROSE_EQ.get(int(mm.group(1)), int(mm.group(1)))})", s)
eq_ref_updates = 0
for p in doc.paragraphs:
    for run in p.runs:
        if re.search(r"Eq\.\s*\(\d+\)", run.text):
            new = _eq_ref_sub(run.text)
            if new != run.text:
                run.text = new
                eq_ref_updates += 1
report.append(f"[OK] prose equation references updated in {eq_ref_updates} runs")

# ---------------------------------------------------------------- 2. FIGURE RENUMBER
FIG_MAP = {7: 4, 8: 5, 9: 6, 10: 7}
def _fig_sub(s):
    return re.sub(r"Figure (7|8|9|10)\b",
                  lambda mm: f"Figure {FIG_MAP[int(mm.group(1))]}", s)
fig_ref_updates = 0
for p in doc.paragraphs:
    for run in p.runs:
        if re.search(r"Figure (7|8|9|10)\b", run.text):
            new = _fig_sub(run.text)
            if new != run.text:
                run.text = new
                fig_ref_updates += 1
report.append(f"[OK] figure references/markers/captions renumbered in {fig_ref_updates} runs")

# ---------------------------------------------------------------- 3. DELETE OLD FIGS 4-6
# remove the drawing runs that embed image4/image5/image6
removed_imgs = []
for p in doc.paragraphs:
    for r in list(p._p.findall(qn("w:r"))):
        blips = r.findall(".//" + qn("a:blip"))
        if not blips:
            continue
        rid = blips[0].get(qn("r:embed"))
        try:
            part = doc.part.rels[rid].target_part.partname
        except Exception:
            part = ""
        if any(str(part).endswith(f"image{n}.png") for n in (4, 5, 6)):
            try:
                doc.part.drop_rel(rid)
            except Exception:
                pass
            r.getparent().remove(r)
            removed_imgs.append(str(part).split("/")[-1])
report.append(f"[OK] removed embedded old figure images: {sorted(removed_imgs)}")

# delete the two orphaned old-figure reference paragraphs
deleted_paras = []
for p in list(doc.paragraphs):
    s = p.text.strip()
    if s.startswith("Figure. 4 shows cumulative") or s.startswith("Figures 5-6 compares"):
        p._p.getparent().remove(p._p)
        deleted_paras.append(s[:40])
report.append(f"[OK] deleted orphaned reference paragraphs: {deleted_paras}")

doc.save(F)

# ---------------------------------------------------------------- 4. RENAME FILES
RENAME = {"Figure7_PoW_energy_sensitivity": "Figure4_PoW_energy_sensitivity",
          "Figure8_PoS_validator_scaling": "Figure5_PoS_validator_scaling",
          "Figure9_computation_vs_communication": "Figure6_computation_vs_communication",
          "Figure10_carbon_gamma_sensitivity": "Figure7_carbon_gamma_sensitivity"}
for old, new in RENAME.items():
    for ext in ("png", "pdf"):
        o = os.path.join(FIGREV, old + "." + ext)
        n = os.path.join(FIGREV, new + "." + ext)
        if os.path.exists(o):
            os.replace(o, n)
report.append(f"[OK] renamed figure files: {list(RENAME.values())}")

print("Saved:", F)
for line in report:
    print("  " + line)
