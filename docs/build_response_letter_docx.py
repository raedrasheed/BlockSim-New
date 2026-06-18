"""Render docs/RESPONSE_TO_REVIEWERS.md to a DOCX for submission convenience.

Lightweight Markdown handling: headings (#/##/###), bold (**...**), blockquotes
(>), bullet (- ), and paragraphs. Sufficient for the response letter.
"""

import os
import re
from docx import Document
from docx.shared import Pt

HERE = os.path.dirname(__file__)
MD = os.path.join(HERE, "RESPONSE_TO_REVIEWERS.md")
OUT = os.path.join(HERE, "RESPONSE_TO_REVIEWERS.docx")


def add_runs(paragraph, text):
    # split on **bold** segments
    for i, seg in enumerate(re.split(r"(\*\*.+?\*\*)", text)):
        if not seg:
            continue
        if seg.startswith("**") and seg.endswith("**"):
            r = paragraph.add_run(seg[2:-2]); r.bold = True
        else:
            paragraph.add_run(seg)


def main():
    doc = Document()
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(10.5)

    for raw in open(MD, encoding="utf-8").read().splitlines():
        line = raw.rstrip()
        if not line.strip():
            continue
        if line.startswith("# "):
            doc.add_heading(line[2:], level=0)
        elif line.startswith("## "):
            doc.add_heading(line[3:], level=1)
        elif line.startswith("### "):
            doc.add_heading(line[4:], level=2)
        elif line.startswith("> "):
            p = doc.add_paragraph()
            r = p.add_run(line[2:]); r.italic = True
        elif line.startswith("- "):
            p = doc.add_paragraph(style="List Bullet")
            add_runs(p, line[2:])
        elif re.match(r"^\d+\.\s", line):
            p = doc.add_paragraph(style="List Number")
            add_runs(p, re.sub(r"^\d+\.\s", "", line))
        elif set(line) <= set("-—"):
            continue  # horizontal rule
        else:
            p = doc.add_paragraph()
            add_runs(p, line)

    doc.save(OUT)
    print("Saved:", OUT)


if __name__ == "__main__":
    main()
