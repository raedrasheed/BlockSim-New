"""Render the revised manuscript to DOCX from the shared content model."""

import os
from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

from manuscript_content import build_blocks, FIG

OUT = os.path.join(os.path.dirname(__file__),
                   "Extending_BlockSim_Consensus_Aware_Energy_Carbon_REVISED.docx")


def main():
    blocks = build_blocks()
    doc = Document()
    st = doc.styles["Normal"]
    st.font.name = "Calibri"
    st.font.size = Pt(11)

    for blk in blocks:
        kind = blk[0]
        if kind == "title":
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run(blk[1]); r.bold = True; r.font.size = Pt(16)
        elif kind == "author":
            p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run(blk[1]); r.bold = blk[3]; r.font.size = Pt(blk[2])
        elif kind == "abstract":
            p = doc.add_paragraph()
            p.add_run("Abstract — ").bold = True
            p.add_run(blk[1])
        elif kind == "keywords":
            p = doc.add_paragraph()
            p.add_run("Keywords: ").bold = True
            p.add_run(blk[1])
        elif kind == "h1":
            doc.add_heading(blk[1], level=1)
        elif kind == "h2":
            doc.add_heading(blk[1], level=2)
        elif kind == "p":
            doc.add_paragraph(blk[1])
        elif kind == "bullet":
            doc.add_paragraph(blk[1], style="List Bullet")
        elif kind == "number":
            doc.add_paragraph(blk[1], style="List Number")
        elif kind == "eq":
            p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run(blk[1]); r.italic = True
            p.add_run("\t\t(" + str(blk[2]) + ")")
        elif kind == "fig":
            _, stem, num, caption = blk
            p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.add_run().add_picture(os.path.join(FIG, stem + ".png"), width=Inches(6.0))
            cap = doc.add_paragraph(); cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
            cr = cap.add_run(f"Figure {num}. {caption}"); cr.italic = True
            cr.font.size = Pt(9.5)
        elif kind == "table":
            _, num, caption, headers, rows = blk
            cap = doc.add_paragraph()
            cr = cap.add_run(f"Table {num}. {caption}"); cr.bold = True
            cr.font.size = Pt(9.5)
            t = doc.add_table(rows=1, cols=len(headers))
            t.style = "Light Grid Accent 1"
            t.alignment = WD_TABLE_ALIGNMENT.CENTER
            hdr = t.rows[0].cells
            for i, h in enumerate(headers):
                run = hdr[i].paragraphs[0].add_run(str(h)); run.bold = True
                run.font.size = Pt(9)
            for row in rows:
                cells = t.add_row().cells
                for i, v in enumerate(row):
                    cells[i].text = str(v)
                    for para in cells[i].paragraphs:
                        for run in para.runs:
                            run.font.size = Pt(9)
            doc.add_paragraph()
        elif kind == "ref":
            p = doc.add_paragraph()
            p.add_run(f"[{blk[1]}] ").bold = True
            p.add_run(blk[2])

    doc.save(OUT)
    print("Saved DOCX:", OUT)
    print("Tables:", len(doc.tables), "Paragraphs:", len(doc.paragraphs))


if __name__ == "__main__":
    main()
